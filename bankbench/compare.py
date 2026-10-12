"""Bandingkan template hasil update dengan workbook yang dikerjakan analis (sel per sel).

Yang dibandingkan: baris yang di salah satu sisi berisi angka input (angka atau rumus angka murni
'=a+b'), di kolom periode yang baru diisi. Baris total/rumus ikut terbanding lewat nilainya kalau
sisi lain menulis angka langsung (mis. TOTAL ASET diketik analis vs '=SUM(...)' di output).
"""
from __future__ import annotations

import io
import re

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .template_model import is_ref_formula, numeric_value

STATUS_SAME, STATUS_DIFF = "Sama", "Beda"
STATUS_ONLY_ANALYST, STATUS_ONLY_OURS = "Hanya di file analis", "Hanya di output"


def _evaluator(ws):
    """Evaluasi rumus sederhana (angka, + - * /, referensi sel, SUM/AVERAGE range) di satu sheet."""
    cache = {}

    def cell(col, row):
        key = (col, row)
        if key in cache:
            return cache[key]
        cache[key] = None  # cegah rekursi siklik
        v = ws.cell(row, col).value
        out = None
        if isinstance(v, bool):
            out = None
        elif isinstance(v, (int, float)):
            out = float(v)
        elif isinstance(v, str) and v.startswith("="):
            out = _eval_formula(v[1:], cell)
        cache[key] = out
        return out

    return cell


def _range_vals(a, b, cell):
    (c1, r1), (c2, r2) = (re.match(r"([A-Z]+)(\d+)", x).groups() for x in (a, b))
    vals = []
    for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
        for r in range(int(r1), int(r2) + 1):
            v = cell(c, r)
            if v is not None:
                vals.append(v)
    return vals


def _eval_formula(expr, cell):
    e = expr.replace("$", "").replace(" ", "")
    if re.search(r"[A-Z]+\(", re.sub(r"(SUM|AVERAGE)\(", "", e)):
        return None  # fungsi lain (IF, VLOOKUP, ...) tidak dievaluasi

    def agg(m):
        vals = []
        for part in m.group(2).split(","):
            if ":" in part:
                vals += _range_vals(*part.split(":"), cell)
            else:
                mm = re.fullmatch(r"([A-Z]+)(\d+)", part)
                v = cell(column_index_from_string(mm.group(1)), int(mm.group(2))) if mm else None
                if v is not None:
                    vals.append(v)
        if m.group(1) == "AVERAGE":
            return repr(sum(vals) / len(vals)) if vals else "0"
        return repr(sum(vals))

    e = re.sub(r"(SUM|AVERAGE)\(([^()]*)\)", agg, e)
    e = re.sub(r"\b([A-Z]{1,3})(\d+)\b",
               lambda m: repr(cell(column_index_from_string(m.group(1)), int(m.group(2))) or 0.0), e)
    if not re.fullmatch(r"[\d.eE+\-*/()]+", e):
        return None
    try:
        return float(eval(e, {"__builtins__": {}}, {}))  # noqa: S307 - hanya angka & operator
    except Exception:
        return None


def _is_input(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) or (
        isinstance(v, str) and v.startswith("=") and not is_ref_formula(v) and numeric_value(v) is not None)


def _label(ws, row, upto_col=6):
    for c in range(upto_col, 0, -1):
        v = ws.cell(row, c).value
        if isinstance(v, str) and re.search(r"[A-Za-z]{2,}", v):
            return v.strip()
    return ""


def find_period_col(ws, label: str, header_row: int | None = None, fallback: int | None = None):
    """Kolom di sheet analis yang judulnya = label periode ('Jun 26'); tidak ketemu -> fallback."""
    norm = re.sub(r"\s+", " ", str(label)).strip().lower()
    rows = [header_row] if header_row else []
    rows += [r for r in range(1, 16) if r != header_row]
    for r in rows:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and re.sub(r"\s+", " ", v).strip().lower() == norm:
                return c
    return fallback


def _close(a, b):
    tol = 0.5 if max(abs(a), abs(b)) >= 1000 else 0.005
    return abs(a - b) <= max(tol, 1e-9 * max(abs(a), abs(b)))


def compare_workbooks(ours_bytes: bytes, analyst_bytes: bytes, targets):
    """targets: list of dict(sheet, col, label, header_row). -> (rows, summary)
    rows: list of dict per sel yang tidak sama; summary: per sheet hitungan status."""
    ours_wb = openpyxl.load_workbook(io.BytesIO(ours_bytes))
    an_wb = openpyxl.load_workbook(io.BytesIO(analyst_bytes))
    an_val = openpyxl.load_workbook(io.BytesIO(analyst_bytes), data_only=True)
    rows, summary = [], []
    for t in targets:
        sheet = t["sheet"]
        if sheet not in an_wb.sheetnames or sheet not in ours_wb.sheetnames:
            summary.append({"Sheet": sheet, "Periode": t["label"], "Keterangan": "sheet tidak ada di file analis"})
            continue
        o_ws, a_ws, a_vs = ours_wb[sheet], an_wb[sheet], an_val[sheet]
        oc = t["col"]
        ac = find_period_col(a_ws, t["label"], t.get("header_row"), fallback=None)
        if ac is None:
            summary.append({"Sheet": sheet, "Periode": t["label"],
                            "Keterangan": f"kolom '{t['label']}' tidak ditemukan di file analis"})
            continue
        o_eval, a_eval = _evaluator(o_ws), _evaluator(a_ws)
        counts = {STATUS_SAME: 0, STATUS_DIFF: 0, STATUS_ONLY_ANALYST: 0, STATUS_ONLY_OURS: 0}
        for r in range(1, max(o_ws.max_row, a_ws.max_row) + 1):
            ov, av = o_ws.cell(r, oc).value, a_ws.cell(r, ac).value
            if not (_is_input(ov) or _is_input(av)):
                continue
            o_num = o_eval(oc, r)
            a_num = a_vs.cell(r, ac).value
            if not isinstance(a_num, (int, float)) or isinstance(a_num, bool):
                a_num = a_eval(ac, r)
            if o_num is None and a_num is None:
                continue
            if o_num is None:
                status = STATUS_ONLY_ANALYST
            elif a_num is None:
                status = STATUS_ONLY_OURS
            else:
                status = STATUS_SAME if _close(o_num, a_num) else STATUS_DIFF
            counts[status] += 1
            if status == STATUS_SAME:
                continue
            note = ""
            prev = a_ws.cell(r, ac - 1).value if ac > 1 else None
            if isinstance(av, str) and _is_input(av) and isinstance(prev, str) and \
                    av.replace(" ", "") == prev.replace(" ", ""):
                note = (f"⚠ rumus di file analis sama persis dengan kolom sebelumnya "
                        f"({a_ws.cell(t.get('header_row') or 1, ac - 1).value or get_column_letter(ac - 1)}) "
                        "— kemungkinan tersalin, cek")
            elif status == STATUS_ONLY_ANALYST:
                note = "tidak terisi di output — biasanya data dari luar laporan publikasi, atau perlu diisi manual"
            rows.append({"Sheet": sheet, "Periode": t["label"], "Baris": r, "Akun": _label(a_ws, r) or _label(o_ws, r),
                         "Output": o_num, "File analis": a_num,
                         "Selisih": (o_num - a_num) if (o_num is not None and a_num is not None) else None,
                         "Status": status, "Isi output": ov, "Isi file analis": av, "Keterangan": note})
        summary.append({"Sheet": sheet, "Periode": t["label"], **counts, "Keterangan": ""})
    return rows, summary
