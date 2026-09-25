"""Identifikasi akun: cocokkan tiap baris template dengan baris laporan publikasi.

Sinyal yang dipakai (digabung jadi satu skor 0-100+):
1. POLA ANGKA  — laporan publikasi selalu memuat angka PEMBANDING (mis. 31 Des
   tahun lalu untuk neraca, 30 Sep tahun lalu untuk laba rugi). Angka pembanding
   itu sudah ada di template (diisi kuartal-kuartal sebelumnya). Kalau satu
   baris laporan punya angka yang PERSIS sama dengan histori satu baris
   template, hampir pasti itu akun yang sama — apa pun namanya.
   Dari sini juga ketahuan otomatis: kolom angka ke berapa = pembanding
   (konsolidasian/bank only), jadi kolom periode berjalan bisa ditebak,
   plus skala (juta/miliar) dan konvensi tanda (+/-).
2. ALIAS       — kamus padanan nama (bawaan + hasil belajar dari koreksi user).
3. NAMA AKUN   — fuzzy similarity label (+ label induk untuk akun kembar
   seperti 'Lainnya', 'Surat berharga', 'Kredit').
4. POSISI      — urutan relatif di dalam seksi (tie-breaker untuk label kembar).
5. KEWAJARAN   — perubahan vs periode sebelumnya (angka yang melonjak 10x
   diturunkan skornya).
Penugasan dilakukan satu-ke-satu (greedy, skor tertinggi dulu) per seksi.
"""
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field

from .aliases import lookup
from .template_model import SheetModel, is_ref_formula
from .text import clean_label, similarity

AUTO_THRESHOLD = 75
_TOTAL_LABEL = re.compile(r"^\s*(total|jumlah|laba \(rugi\)|laba bersih|pendapatan \(beban\) .* bersih)", re.I)
MIN_ANCHOR_ABS = 1000
RELEVANT_SECTIONS = ("BS", "IS", "CAP")


@dataclass
class MatchResult:
    row: int
    section: str
    label: str
    parent: str | None
    kind: str  # 'Nilai' | 'Formula' | 'Biasanya kosong'
    prev_value: float | None = None
    source_idx: int | None = None
    source_label: str | None = None
    value: float | None = None
    score: float = 0.0
    method: str = "-"
    note: str = ""


@dataclass
class Layout:
    n_values: int
    comp_idx: int | None
    cur_idx: int
    scale: float
    anchors: int
    comp_period_by_section: dict = field(default_factory=dict)
    entity: str | None = None  # 'KONS' / 'IND' — dari header tabel
    use_header: bool = False  # True -> kolom dipilih per baris dari header (periode+entitas)


def _close(a: float, b: float) -> bool:
    return abs(a - b) <= max(1.0, 1e-6 * abs(b))


_EXPENSE = re.compile(r"\b(beban|expenses?|impairment|kerugian terkait|losses related|taksiran pajak|estimated current)",
                      re.I)
_NOT_EXPENSE = re.compile(r"keuntungan|gain|\(beban\)|\(expenses?\)|pendapatan|income|revenue|bersih|net\b", re.I)


def _is_expense(*labels) -> bool:
    txt = " ".join(x or "" for x in labels)
    return bool(_EXPENSE.search(txt)) and not _NOT_EXPENSE.search(labels[0] or "")


class Matcher:
    def __init__(self, model: SheetModel, source_rows, target_col: int, aliases: dict, period=None):
        """period: (bulan, tahun) laporan — dipakai untuk memilih kolom angka dari
        header tabel ('31 Des 2025'). Default: dari judul kolom tujuan di template."""
        self.model = model
        from .periods import parse_period_label
        self.period = period or parse_period_label(model.period_cols.get(target_col))
        self.source = [s for s in source_rows if s.section in RELEVANT_SECTIONS + ("?",)]
        self.target_col = target_col
        self.aliases = aliases
        self.ref_col = model.reference_col(target_col)
        self.hist_cols = model.filled_cols_before(target_col)[-8:]
        self._value_rows = []
        self._formula_rows = set()
        for tr in model.rows:
            if tr.is_formula_at(self.ref_col) or tr.is_formula_at(target_col):
                self._formula_rows.add(tr.row)
            else:
                self._value_rows.append(tr)
        # Konvensi template: kadang angka akun ditaruh di SUB-baris, bukan di baris
        # induknya (mis. 'Aset antar kantor' kosong, angkanya di 'a. Melakukan kegiatan
        # operasional di Indonesia'; 'Kredit' kosong, angkanya di 'd. Pinjaman yang
        # diberikan dan piutang'). Baris anak itu ikut "mewarisi" nama induknya.
        rows_by_num = {r.row: r for r in model.rows}
        filled_children = defaultdict(list)
        for tr in self._value_rows:
            if tr.parent_row and self.usually_filled(tr):
                filled_children[tr.parent_row].append(tr)
        self._redirect, self._redirected_parents = {}, {}
        for prow, kids in filled_children.items():
            ptr = rows_by_num.get(prow)
            if ptr and ptr.row not in self._formula_rows and not self.usually_filled(ptr) and len(kids) == 1:
                self._redirect[kids[0].row] = ptr
                self._redirected_parents[prow] = kids[0].row
        sec_rows = defaultdict(list)
        for s in self.source:
            sec_rows[s.section].append(s)
        self._src_pos = {}
        for sec, rows in sec_rows.items():
            for i, s in enumerate(rows):
                self._src_pos[s.idx] = i / max(1, len(rows) - 1)

    # ------------------------------------------------------------------ util
    def history(self, tr, n=8):
        out = []
        for c in self.hist_cols[-n:]:
            v = tr.numeric(c)
            if v is not None:
                out.append((c, v))
        return out

    def last_value(self, tr):
        h = self.history(tr)
        return h[-1][1] if h else None

    def usually_filled(self, tr) -> bool:
        return any(tr.numeric(c) is not None for c in self.hist_cols[-4:])

    def _tpl_pos(self, tr):
        start, end = self.model.sections[tr.section]
        return (tr.row - start) / max(1, end - start)

    @staticmethod
    def _sections_ok(tr, sr):
        return sr.section == "?" or sr.section == tr.section

    # ------------------------------------------------------------ layout
    def _anchor_index(self, scale):
        idx = defaultdict(list)
        for s in self.source:
            for k, v in enumerate(s.values):
                a = abs(v * scale)
                if a >= MIN_ANCHOR_ABS:
                    idx[round(a)].append((s, k))
        return idx

    def _find_anchors(self, scale):
        """-> list of (tr, sr, k, hist_col, sign) : angka laporan == histori template."""
        idx = self._anchor_index(scale)
        found = []
        for tr in self._value_rows:
            seen = set()
            for c, h in self.history(tr):
                if abs(h) < MIN_ANCHOR_ABS:
                    continue
                key = round(abs(h))
                for kk in (key - 1, key, key + 1):
                    for s, k in idx.get(kk, ()):
                        if not self._sections_ok(tr, s) or (s.idx, k) in seen:
                            continue
                        v = s.values[k] * scale
                        if _close(abs(v), abs(h)):
                            seen.add((s.idx, k))
                            found.append((tr, s, k, c, 1 if (v >= 0) == (h >= 0) else -1))
        return found

    def detect_layout(self, cur_idx_override=None) -> Layout:
        counts = Counter(len(s.values) for s in self.source if s.section in RELEVANT_SECTIONS) or Counter(
            len(s.values) for s in self.source)
        n_values = counts.most_common(1)[0][0] if counts else 1
        best = None
        for scale in (1.0, 0.001, 1000.0):
            anchors = self._find_anchors(scale)
            if best is None or len(anchors) > len(best[1]):
                best = (scale, anchors)
        scale, anchors = best
        self._anchors = anchors
        comp_idx = None
        comp_by_sec = {}
        # jumlah akun (unik) yang cocok per posisi angka; harus cukup banyak supaya
        # tidak ketipu kebetulan (mis. 'Modal dasar' yang nilainya tidak pernah berubah)
        rows_per_k = defaultdict(set)
        for tr, _, k, _, _ in anchors:
            rows_per_k[k].add(tr.row)
        n_filled = sum(1 for tr in self._value_rows if self.usually_filled(tr))
        min_needed = max(8, int(0.15 * n_filled))
        if rows_per_k:
            k_best = max(rows_per_k, key=lambda k: len(rows_per_k[k]))
            if len(rows_per_k[k_best]) >= min_needed:
                comp_idx = k_best
        if comp_idx is not None:
            per_sec = defaultdict(Counter)
            for tr, _, k, c, _ in anchors:
                if k == comp_idx:
                    per_sec[tr.section][c] += 1
            comp_by_sec = {sec: self.model.period_cols.get(cnt.most_common(1)[0][0]) for sec, cnt in per_sec.items()}
        if cur_idx_override is not None:
            cur_idx = cur_idx_override
        elif comp_idx is not None:
            cur_idx = comp_idx - 1 if comp_idx >= 1 else min(1, n_values - 1)
        else:
            # tanpa histori: tebak dari susunan kolom umum laporan OJK
            # (Bank kini | Bank lalu | Konsolidasian kini | Konsolidasian lalu)
            cur_idx = 2 if n_values >= 4 else 0
        if comp_idx is None:
            scale = self._guess_scale(cur_idx)
        # Header tabel ('INDIVIDUAL | KONSOLIDASIAN' + '31 Des 2025 | 31 Des 2024')
        # -> entitas yang dipakai template = entitas kolom pembanding yang cocok histori
        entity, use_header = None, False
        with_hdr = [s for s in self.source if s.cols and s.section in RELEVANT_SECTIONS]
        if cur_idx_override is None and self.period and len(with_hdr) >= 0.5 * max(1, len(
                [s for s in self.source if s.section in RELEVANT_SECTIONS])):
            ent_votes = Counter()
            for tr, s, k, _, _ in anchors:
                meta = self._col_meta(s)
                if meta and k < len(meta) and meta[k][1]:
                    ent_votes[meta[k][1]] += 1
            if ent_votes:
                entity = ent_votes.most_common(1)[0][0]
            elif any(e == "KONS" for s in with_hdr for _, e in s.cols):
                entity = "KONS"
            use_header = any(p == self.period for s in with_hdr for p, _ in s.cols)
        self.layout = Layout(n_values, comp_idx, cur_idx, scale, len(anchors), comp_by_sec, entity, use_header)
        return self.layout

    @staticmethod
    def _col_meta(sr):
        """Metadata kolom sejajar dengan sr.values (kalau angkanya lebih sedikit dari
        kolom header, diasumsikan rata kanan — mis. baris yang hanya punya angka
        Konsolidasian)."""
        if not sr.cols:
            return None
        if len(sr.values) == len(sr.cols):
            return sr.cols
        if len(sr.values) < len(sr.cols):
            return sr.cols[len(sr.cols) - len(sr.values):]
        return None

    def _index_for(self, sr):
        lay = self.layout
        if lay.use_header:
            meta = self._col_meta(sr)
            if meta:
                hits = [i for i, (p, e) in enumerate(meta) if p == self.period and (lay.entity is None or e in (lay.entity, None))]
                if hits:
                    return hits[-1] if lay.entity is None else hits[0]
                if any(p for p, _ in meta):
                    return None  # header ada tapi periode laporan tidak ada di tabel ini
        i = lay.cur_idx
        # baris yang hanya berisi separuh kolom (mis. hanya Konsolidasian)
        if len(sr.values) <= i and lay.n_values >= 4 and len(sr.values) == lay.n_values // 2 and i >= lay.n_values // 2:
            return i - lay.n_values // 2
        return i if i < len(sr.values) else None

    def _guess_scale(self, cur_idx):
        """Tanpa pola angka: bandingkan akun yang namanya persis sama dengan nilai
        periode lalu di template -> ketahuan laporan dalam juta / miliar / ribuan."""
        ratios = []
        for tr in self._value_rows:
            last = self.last_value(tr)
            if not last or abs(last) < MIN_ANCHOR_ABS:
                continue
            for sr in self.source:
                if self._sections_ok(tr, sr) and similarity(tr.label, sr.label) >= 97:
                    v = sr.value_at(cur_idx)
                    if v:
                        ratios.append(abs(v / last))
                    break
        if len(ratios) < 5:
            return 1.0
        ratios.sort()
        med = ratios[len(ratios) // 2]
        if med > 200:
            return 0.001
        if med < 0.005:
            return 1000.0
        return 1.0

    # ------------------------------------------------------------ scoring
    def current_value(self, sr):
        i = self._index_for(sr)
        v = None if i is None else sr.value_at(i)
        return None if v is None else v * self.layout.scale

    def _label_score(self, tr, sr):
        s = similarity(tr.label, sr.label)
        if tr.parent and sr.parent:
            s = 0.75 * s + 0.25 * similarity(tr.parent, sr.parent)
        ptr = self._redirect.get(tr.row)
        if ptr is not None:
            s = max(s, 0.97 * self._label_score(ptr, sr))
        return s

    def _alias_score(self, tr, sr):
        ptr = self._redirect.get(tr.row)
        if ptr is not None:
            inherited = self._alias_score(ptr, sr)
            if inherited:
                return inherited
        names = lookup(self.aliases, tr.section, tr.label, tr.parent)
        if not names or clean_label(sr.label) not in names:
            return 0.0
        if tr.parent and sr.parent and similarity(tr.parent, sr.parent) < 55:
            # alias tanpa induk ('surat berharga') jangan nyasar ke anak 'Cadangan kerugian'
            specific = lookup({tr.section: {k: v for k, v in self.aliases.get(tr.section, {}).items() if " | " in k}},
                              tr.section, tr.label, tr.parent)
            if clean_label(sr.label) not in specific:
                return 0.0
        return 97.0

    def _plausibility(self, tr, v):
        last = self.last_value(tr)
        if v is None or not last:
            return 0.0, ""
        ratio = v / last
        if ratio < 0:
            return -4.0, "tanda berbeda dari periode lalu"
        r = abs(ratio)
        if 0.5 <= r <= 2:
            return 4.0, ""
        if r > 10 or r < 0.1:
            return -12.0, f"berubah {r:.1f}x vs periode lalu"
        return 0.0, ""

    def score_pairs(self):
        anchor_map = defaultdict(dict)  # tr.row -> sr.idx -> sign
        for tr, s, k, _, sign in self._anchors:
            if k != self._index_for(s):
                anchor_map[tr.row][s.idx] = sign
        pairs = []
        for tr in self._value_rows:
            if tr.row in self._redirected_parents:
                continue
            tpos = self._tpl_pos(tr)
            for sr in self.source:
                if not self._sections_ok(tr, sr):
                    continue
                v = self.current_value(sr)
                label_s = self._label_score(tr, sr)
                alias_s = self._alias_score(tr, sr)
                anchored = sr.idx in anchor_map.get(tr.row, {})
                if anchored:
                    base, method = 90 + 0.1 * max(label_s, alias_s), "Pola angka"
                elif alias_s >= label_s and alias_s > 0:
                    base, method = alias_s, "Alias"
                else:
                    base, method = label_s, "Nama akun"
                if base < 55:
                    continue
                pos_bonus = 6 * max(0.0, 1 - abs(tpos - self._src_pos.get(sr.idx, 0.5)) / 0.25)
                plaus, note = self._plausibility(tr, v)
                if anchored:
                    plaus, note = max(plaus, 0.0), ""
                penalty = -5 if sr.section == "?" else 0
                pairs.append((base + pos_bonus + plaus + penalty, tr, sr, method, note, anchor_map[tr.row].get(sr.idx)))
        pairs.sort(key=lambda p: p[0], reverse=True)
        return pairs

    def value_for(self, tr, sr, sign=None):
        """Nilai periode berjalan dari baris laporan `sr` untuk baris template `tr`
        (sudah dikali skala & disesuaikan konvensi tanda). -> (nilai, catatan[])"""
        notes = []
        v = self.current_value(sr)
        if v is None:
            return None, ["kolom angka tidak tersedia di baris ini"]
        hist = [h for _, h in self.history(tr, 4) if h != 0]
        # akun pengurang ('-/-', mis. CKPN, akumulasi penyusutan) kadang ditulis
        # positif di laporan, padahal di template selalu minus
        contra = any("-/-" in (x or "") for x in (tr.label, tr.parent, sr.label, sr.parent))
        if sign == -1 or (sign is None and contra and v > 0 and len(hist) >= 2 and all(h < 0 for h in hist)):
            v = -v
            notes.append("tanda disesuaikan ke konvensi template")
        # akun BEBAN: laporan kadang menulis beban dalam kurung (negatif), template
        # menyimpan positif (atau sebaliknya) -> ikuti tanda histori template
        if v and len(hist) >= 2 and _is_expense(tr.label, sr.label):
            want = 1 if all(h > 0 for h in hist) else -1 if all(h < 0 for h in hist) else 0
            if want and (v > 0) != (want > 0):
                v = -v
                if "tanda disesuaikan ke konvensi template" not in notes:
                    notes.append("tanda disesuaikan ke konvensi template")
        return v, notes

    # ------------------------------------------------------------ main
    def run(self, cur_idx_override=None):
        if not hasattr(self, "layout") or cur_idx_override is not None:
            self.detect_layout(cur_idx_override)
        assigned, used = {}, set()
        for score, tr, sr, method, note, sign in self.score_pairs():
            if tr.row in assigned or sr.idx in used:
                continue
            assigned[tr.row] = (score, sr, method, note, sign)
            used.add(sr.idx)
        self._fix_crossings(assigned)
        self.extras = self._find_sum_patterns(assigned)

        results = []
        for tr in self.model.rows:
            prev = self.last_value(tr)
            if tr.row in self._formula_rows:
                results.append(MatchResult(tr.row, tr.section, tr.label, tr.parent, "Formula", prev))
                continue
            filled = self.usually_filled(tr)
            res = MatchResult(tr.row, tr.section, tr.label, tr.parent, "Nilai" if filled else "Biasanya kosong", prev)
            if tr.row in self._redirected_parents:
                res.note = f"angka akun ini biasa ditaruh di baris {self._redirected_parents[tr.row]}"
            if tr.row in assigned:
                score, sr, method, note, sign = assigned[tr.row]
                v, notes = self.value_for(tr, sr, sign)
                if note:
                    notes.insert(0, note)
                if tr.row in self.extras and v is not None:
                    base = self.current_value(sr)
                    factor = (v / base) if base else 1.0
                    for ex in self.extras[tr.row]:
                        v += factor * (self.current_value(ex) or 0.0)
                    notes.insert(0, "+ " + " + ".join(f"'{ex.label}'" for ex in self.extras[tr.row])
                                 + " (digabung sesuai pola histori template)")
                    method = "Pola penjumlahan"
                res.source_idx, res.source_label = sr.idx, sr.label
                res.score, res.method = round(min(score, 100.0), 1), method
                # baris yang di histori template biasanya kosong (induk/subtotal, akun yang
                # ditaruh di baris lain) TIDAK diisi otomatis -> cuma saran, hindari dobel hitung
                ok = filled and score >= AUTO_THRESHOLD and v is not None
                if ok:
                    res.value = v
                elif filled:
                    notes.append("skor rendah — cek manual")
                elif v not in (None, 0):
                    notes.append("biasanya kosong di template — tidak diisi otomatis (saran saja)")
                res.note = "; ".join(n for n in notes if n)
            elif tr.row in getattr(self, "combo_rows", {}):
                parts = self.combo_rows[tr.row]
                cur = [self.current_value(x) for x in parts]
                if all(v is not None for v in cur):
                    res.value = sum(cur)
                res.source_idx, res.source_label = parts[0].idx, parts[0].label
                res.score, res.method = 95.0, "Pola penjumlahan"
                res.note = "= " + " + ".join(f"'{x.label}'" for x in parts) + " (sesuai pola histori template)"
            elif filled:
                res.note = "tidak ketemu di laporan — isi manual"
            results.append(res)
        return results

    def comp_value(self, sr, section):
        """Angka PEMBANDING (periode lalu) dari baris laporan untuk seksi tsb."""
        from .periods import parse_period_label
        lay = self.layout
        label = lay.comp_period_by_section.get(section)
        per = parse_period_label(label) if label else None
        meta = self._col_meta(sr)
        if lay.use_header and meta and per:
            for i, (p, e) in enumerate(meta):
                if p == per and (lay.entity is None or e in (lay.entity, None)):
                    return sr.values[i] * lay.scale
            return None
        if lay.comp_idx is None:
            return None
        v = sr.value_at(lay.comp_idx)
        return None if v is None else v * lay.scale

    def _find_sum_patterns(self, assigned):
        """Konvensi analis: satu baris template kadang = JUMLAH beberapa akun laporan
        (mis. 'Aset lainnya' = Aset lainnya + Aset keuangan lainnya). Terdeteksi dari
        histori: angka pembanding laporan != histori template, tapi selisihnya persis
        sama dengan 1-3 akun laporan lain yang belum terpakai.
        Baris template yang belum punya pasangan sama sekali juga dicoba: mungkin
        di laporan dipecah (mis. 'Penghasilan komprehensif lain' = Keuntungan + Kerugian).
        Return: row -> list akun tambahan. Baris tanpa pasangan yang ketemu kombinasinya
        disimpan di self.combo_rows (row -> list akun)."""
        import itertools

        self.combo_rows = {}
        if not getattr(self, "layout", None) or not self.layout.comp_period_by_section:
            return {}
        rows_by_num = {r.row: r for r in self.model.rows}
        label_to_col = {v: c for c, v in self.model.period_cols.items()}
        used = {a[1].idx for r, a in assigned.items() if self.usually_filled(rows_by_num[r])}
        extras = {}

        def hist_at_comp(tr):
            ccol = label_to_col.get(self.layout.comp_period_by_section.get(tr.section))
            return tr.numeric(ccol) if ccol else None

        def candidates(tr):
            out = []
            for s in self.source:
                if s.idx in used or not self._sections_ok(tr, s) or s.section == "?":
                    continue
                cv = self.comp_value(s, tr.section)
                if cv is not None and abs(cv) >= 1 and s.values and not _TOTAL_LABEL.search(s.label):
                    out.append((s, cv))
            return out[:60]

        def unique_combo(cands, target, max_k=3):
            for k in range(1, max_k + 1):
                hits = [c for c in itertools.combinations(cands, k) if _close(sum(cv for _, cv in c), target)]
                if len(hits) == 1:
                    return [s for s, _ in hits[0]]
                if len(hits) > 1:
                    return None
            return None

        for row, (score, sr, method, note, sign) in sorted(assigned.items()):
            tr = rows_by_num[row]
            if method == "Pola angka" or not self.usually_filled(tr):
                continue
            h, c = hist_at_comp(tr), self.comp_value(sr, tr.section)
            if h is None or c is None:
                continue
            d = h - (-c if sign == -1 else c)
            if abs(d) < MIN_ANCHOR_ABS:
                continue
            combo = unique_combo(candidates(tr), d)
            if combo:
                extras[row] = combo
                used.update(x.idx for x in combo)

        for tr in self._value_rows:
            if tr.row in assigned or tr.row in self._redirected_parents or not self.usually_filled(tr):
                continue
            h = hist_at_comp(tr)
            if h is None or abs(h) < MIN_ANCHOR_ABS:
                continue
            # minimal salah satu akun (atau induknya) namanya mirip -> hindari kebetulan angka
            cands = [(s, cv) for s, cv in candidates(tr)
                     if max(similarity(tr.label, s.label), similarity(tr.label, s.parent or "")) >= 60]
            combo = unique_combo(cands, h)
            if combo and len(combo) >= 2:
                self.combo_rows[tr.row] = combo
                used.update(x.idx for x in combo)
        return extras

    def _fix_crossings(self, assigned):
        """Label kembar (mis. 'PEMILIK', 'KEPENTINGAN NON PENGENDALI' muncul 2x) harus
        dipasangkan sesuai URUTAN: kalau dua pasangan saling 'menyilang' dan
        namanya sama-sama cocok, tukar."""
        rows_by_num = {r.row: r for r in self.model.rows}
        changed = True
        guard = 0
        while changed and guard < 50:
            changed, guard = False, guard + 1
            items = sorted(assigned.items())
            for i in range(len(items)):
                for j in range(i + 1, len(items)):
                    (r1, a1), (r2, a2) = items[i], items[j]
                    t1, t2 = rows_by_num[r1], rows_by_num[r2]
                    s1, s2 = a1[1], a2[1]
                    if t1.section != t2.section or s1.idx < s2.idx:
                        continue
                    if a1[2] == "Pola angka" or a2[2] == "Pola angka":
                        continue
                    if clean_label(s1.label) != clean_label(s2.label):
                        continue
                    assigned[r1], assigned[r2] = (a2[0], s2, a2[2], a2[3], a2[4]), (a1[0], s1, a1[2], a1[3], a1[4])
                    changed = True
                    break
                if changed:
                    break

    # ------------------------------------------------------------ checks
    def label_candidates(self, tr, min_score=85):
        """Kandidat baris laporan (skor, SourceRow) yang namanya mirip, urut skor."""
        out = []
        tpos = self._tpl_pos(tr)
        for sr in self.source:
            if not self._sections_ok(tr, sr):
                continue
            s = max(self._label_score(tr, sr), self._alias_score(tr, sr))
            if s >= min_score:
                s += 6 * max(0.0, 1 - abs(tpos - self._src_pos.get(sr.idx, 0.5)) / 0.25)
                out.append((s, sr))
        out.sort(key=lambda x: x[0], reverse=True)
        return out

    def best_label_match(self, tr, min_score=85):
        best = None
        tpos = self._tpl_pos(tr)
        for sr in self.source:
            if not self._sections_ok(tr, sr):
                continue
            s = max(self._label_score(tr, sr), self._alias_score(tr, sr))
            if s < min_score:
                continue
            s += 6 * max(0.0, 1 - abs(tpos - self._src_pos.get(sr.idx, 0.5)) / 0.25)
            if best is None or s > best[0]:
                best = (s, sr)
        return best


# ---------------------------------------------------------------------------------
# Evaluasi formula sederhana (SUM, +, -, *, /) untuk rekonsiliasi total
# ---------------------------------------------------------------------------------
_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_SUM = re.compile(r"SUM\(([^()]*)\)", re.I)


def _col_idx(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def translate_formula(formula: str, src_col: int, dst_col: int, row: int) -> str:
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter

    return Translator(formula, origin=f"{get_column_letter(src_col)}{row}").translate_formula(
        f"{get_column_letter(dst_col)}{row}")


def evaluate_column(model: SheetModel, target_col: int, ref_col: int, values: dict):
    """Hitung nilai formula (yang sederhana) di kolom tujuan memakai nilai baru.
    values: row -> float. Return row -> float|None untuk semua baris formula."""
    rows_by_num = {r.row: r for r in model.rows}
    cache = {}

    def formula_of(r):
        tr = rows_by_num.get(r)
        if tr is None:
            return None
        f = tr.cells.get(target_col)
        if is_ref_formula(f):
            return f
        f = tr.cells.get(ref_col)
        if is_ref_formula(f):
            return translate_formula(f, ref_col, target_col, r)
        return None

    def val(c, r, depth=0):
        if c != target_col:
            tr = rows_by_num.get(r)
            v = tr.numeric(c) if tr else None
            return 0.0 if v is None else v
        if r in values and values[r] is not None:
            return float(values[r])
        if r in cache:
            return cache[r]
        f = formula_of(r)
        if f is None or depth > 40:
            tr = rows_by_num.get(r)
            v = tr.numeric(target_col) if tr else None
            return 0.0 if v is None else v
        cache[r] = ev(f, depth + 1)
        return cache[r]

    def ev(f, depth):
        expr = f.strip().lstrip("=")

        def sum_repl(m):
            parts = []
            for arg in m.group(1).split(","):
                arg = arg.strip()
                if ":" in arg:
                    a, b = arg.split(":")
                    ma, mb = _REF.fullmatch(a.strip()), _REF.fullmatch(b.strip())
                    if not ma or not mb or ma.group(1) != mb.group(1):
                        raise ValueError("range")
                    c = ma.group(1)
                    for rr in range(int(ma.group(2)), int(mb.group(2)) + 1):
                        parts.append(f"{c}{rr}")
                else:
                    parts.append(arg)
            return "(" + "+".join(parts or ["0"]) + ")"

        expr = _SUM.sub(sum_repl, expr)
        if re.search(r"[A-Z]{2,}\(", expr, re.I):
            return None  # IF/AVERAGE dll -> tidak dievaluasi
        py = _REF.sub(lambda m: f"_v({_col_idx(m.group(1))},{m.group(2)})", expr)
        if re.search(r"[^\d\s+\-*/().,_v]", py):
            return None
        try:
            out = eval(py, {"__builtins__": {}}, {"_v": lambda c, r: val(c, r, depth)})  # noqa: S307
            return None if out is None or (isinstance(out, float) and math.isnan(out)) else float(out)
        except Exception:
            return None

    out = {}
    for tr in model.rows:
        if formula_of(tr.row):
            out[tr.row] = val(target_col, tr.row)
    return out


_CHECK_LABEL = re.compile(r"total|jumlah|laba|pendapatan bunga bersih|modal inti|modal pelengkap", re.I)


def reconcile(matcher: Matcher, values: dict):
    """Bandingkan total hasil hitung formula template vs total di laporan."""
    model = matcher.model
    computed = evaluate_column(model, matcher.target_col, matcher.ref_col, values)
    rows_by_num = {r.row: r for r in model.rows}
    checks = []
    for r, v in computed.items():
        tr = rows_by_num[r]
        if not _CHECK_LABEL.search(tr.label):
            continue
        if v is None:
            continue
        # label total kadang kembar ('TOTAL LABA TAHUN BERJALAN' 2x) -> kalau salah satu
        # kandidat yang namanya mirip angkanya cocok, anggap OK
        cands = matcher.label_candidates(tr)
        if not cands:
            continue
        exact = [sr for _, sr in cands if matcher.current_value(sr) is not None
                 and abs(abs(v) - abs(matcher.current_value(sr))) <= 5]
        best_sr = exact[0] if exact else cands[0][1]
        m = (None, best_sr)
        src = matcher.current_value(best_sr)
        if src is None:
            continue
        diff = abs(v) - abs(src)
        ok = abs(diff) <= 5  # toleransi pembulatan (juta rupiah)
        info = ""
        if not ok:
            # selisih persis = satu akun laporan -> biasanya beda klasifikasi
            same = [s for s in matcher.source if s.section == tr.section and matcher.current_value(s) is not None
                    and abs(abs(matcher.current_value(s)) - abs(diff)) <= 1 and abs(diff) >= 1]
            if same:
                info = f"selisih = '{same[0].label}' (cek klasifikasi)"
        checks.append({
            "Seksi": tr.section,
            "Baris": r,
            "Total (template)": tr.label,
            "Hasil hitung": v,
            "Di laporan": src,
            "Akun laporan": m[1].label,
            "Selisih": diff,
            "Status": "OK" if ok else "BEDA",
            "Keterangan": info,
        })
    # neraca harus seimbang
    ta = next((r for r in model.rows if r.section == "BS" and clean_label(r.label) == "total aset"), None)
    tl = next((r for r in model.rows if r.section == "BS" and re.search(r"jumlah (kewajiban|liabilitas) dan (modal|ekuitas)",
                                                                          clean_label(r.label))), None)
    if ta and tl and computed.get(ta.row) is not None and computed.get(tl.row) is not None:
        d = computed[ta.row] - computed[tl.row]
        checks.append({
            "Seksi": "BS", "Baris": tl.row, "Total (template)": "Aset = Liabilitas + Ekuitas",
            "Hasil hitung": computed[tl.row], "Di laporan": computed[ta.row], "Akun laporan": "TOTAL ASET (template)",
            "Selisih": d, "Status": "OK" if abs(d) <= 5 else "BEDA", "Keterangan": "",
        })
    return checks
