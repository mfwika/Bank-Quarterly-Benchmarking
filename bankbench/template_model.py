"""Baca struktur sheet per-bank di template.

Sheet bank di template punya pola yang sama:
  - baris header periode (mis. baris 6: 'Mar 23', 'Jun 23', ... 'Dec 25')
  - seksi NERACA (Balance Sheet)            -> 'BS'
  - seksi PERHITUNGAN LABA RUGI (Income St.) -> 'IS'
  - seksi KOMITMEN & KONTINJENSI             -> diabaikan
  - seksi CAPITAL ADEQUACY RATIO (Struktur Modal / KPMM) -> 'CAP'
  - baris-baris rasio/turunan di bawahnya (formula) -> tidak diisi, hanya
    formulanya yang diperpanjang kalau kolom baru ditambahkan.

Semua dibaca dari "grid" (list baris berisi nilai sel) supaya bisa dipakai
baik dari workbook read-only (cepat) maupun workbook biasa.
"""
import re
from dataclasses import dataclass, field

from .periods import looks_like_period, parse_period_label
from .text import enumerator_kind, leading_enumerator, strip_leading_enumerators

SECTION_NAMES = {
    "BS": "Neraca / Balance Sheet",
    "IS": "Laba Rugi / Income Statement",
    "CAP": "Struktur Modal / KPMM",
}

MAX_COLS = 150
_FORMULA_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+")
_PURE_NUMERIC_FORMULA = re.compile(r"^=\s*[+\-]?[\d.\s+\-*/()]+$")

_HEAD_BS = re.compile(r"NERACA|POSISI KEUANGAN|BALANCE SHEET")
_HEAD_IS = re.compile(r"LABA\s*RUGI|LABA/RUGI|INCOME STATEMENT|PROFIT AND LOSS")
_HEAD_COMMIT = re.compile(r"KOMITMEN|KONTI[NJ]")
_HEAD_CAP = re.compile(r"CAPITAL ADEQUACY|PENYEDIAAN MODAL MINIMUM|KPMM")
_CAP_LAST_ROW = re.compile(r"rasio kewajiban penyediaan modal|rasio kpmm", re.I)

_SKIP_LABELS = {"no", "pos-pos", "pos - pos", "in million rp", "in million rp / percentage (%)"}


def numeric_value(v):
    """Nilai sel -> float kalau berupa angka / formula angka murni ('=906137+27063')."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and _PURE_NUMERIC_FORMULA.match(v.strip()):
        try:
            return float(eval(v.strip()[1:], {"__builtins__": {}}, {}))  # noqa: S307 - hanya angka & operator
        except Exception:
            return None
    return None


def is_ref_formula(v) -> bool:
    return isinstance(v, str) and v.strip().startswith("=") and bool(_FORMULA_REF.search(v))


def _is_heading(text: str) -> bool:
    # 'NERACA - Consol' / 'CAPITAL ADEQUACY RATIO - Consol' -> cek bagian utamanya saja
    text = re.split(r"\s[-–]\s|\(", text)[0]
    letters = [c for c in text if c.isalpha()]
    if len(letters) < 6:
        return False
    upper = sum(1 for c in letters if c.isupper())
    return upper / len(letters) > 0.8


@dataclass
class TemplateRow:
    row: int
    label: str
    section: str
    level: int
    parent: str | None = None
    parent_row: int | None = None
    cells: dict = field(default_factory=dict)  # col -> nilai mentah di kolom periode

    def numeric(self, col):
        return numeric_value(self.cells.get(col))

    def is_formula_at(self, col) -> bool:
        return is_ref_formula(self.cells.get(col))


@dataclass
class SheetModel:
    name: str
    bank_name: str
    header_row: int
    period_cols: dict  # col -> label header
    sections: dict  # 'BS' -> (start_row, end_row)
    rows: list  # list[TemplateRow] (hanya baris di dalam seksi)
    max_row: int
    special_rows: dict = field(default_factory=dict)  # 'annualizer'/'period_date' -> row

    # -- periode ------------------------------------------------------------------
    def parsed_periods(self):
        out = {}
        for c, lbl in self.period_cols.items():
            p = parse_period_label(lbl)
            if p:
                out[c] = p
        return out

    def last_period_col(self) -> int:
        return max(self.period_cols)

    def find_period_col(self, month: int, year: int):
        for c, p in self.parsed_periods().items():
            if p == (month, year):
                return c
        return None

    def target_col_for(self, month: int, year: int):
        """(kolom tujuan, sudah_ada?). Kalau periodenya belum ada -> kolom baru
        tepat setelah kolom periode terakhir."""
        c = self.find_period_col(month, year)
        if c is not None:
            return c, True
        return self.last_period_col() + 1, False

    def fill_ratio(self, col) -> float:
        vals = [r for r in self.rows if not r.is_formula_at(col)]
        if not vals:
            return 0.0
        return sum(1 for r in vals if r.numeric(col) is not None) / len(vals)

    def filled_cols_before(self, target_col: int):
        return [c for c in sorted(self.period_cols) if c < target_col and self.fill_ratio(c) > 0.05]

    def reference_col(self, target_col: int):
        """Kolom 'periode sebelumnya' yang datanya sudah terisi — acuan format &
        formula untuk kolom tujuan."""
        cols = self.filled_cols_before(target_col)
        if cols:
            return cols[-1]
        before = [c for c in sorted(self.period_cols) if c < target_col]
        return before[-1] if before else min(self.period_cols)

    def latest_period(self):
        """Periode terakhir yang sudah ada datanya (month, year)."""
        parsed = self.parsed_periods()
        cols = [c for c in self.filled_cols_before(10**6) if c in parsed]
        return parsed[cols[-1]] if cols else None


# ---------------------------------------------------------------------------------
def grid_from_worksheet(ws, max_cols=MAX_COLS):
    return [list(r) for r in ws.iter_rows(max_col=max_cols, values_only=True)]


def _cell(grid, r, c):
    """Akses 1-based yang aman."""
    if r < 1 or r > len(grid):
        return None
    row = grid[r - 1]
    if c < 1 or c > len(row):
        return None
    return row[c - 1]


def find_header_row(grid, max_scan_rows=40, max_scan_cols=MAX_COLS):
    best_row, best_count = None, 0
    for r in range(1, min(max_scan_rows, len(grid)) + 1):
        count = sum(1 for c in range(1, max_scan_cols + 1) if looks_like_period(_cell(grid, r, c)))
        if count > best_count:
            best_row, best_count = r, count
    return best_row, best_count


def is_bank_sheet(first_rows) -> bool:
    txt = " ".join(str(v) for row in first_rows for v in row if isinstance(v, str)).upper()
    return bool(_HEAD_BS.search(txt))


def _row_texts(grid, r, label_zone_end, raw=False):
    """-> (label, label_col, enumerator_kind) untuk satu baris
    (+ token enumerator mentah kalau raw=True)."""
    label, label_col, enum, token = None, None, None, None
    for c in range(1, label_zone_end):
        v = _cell(grid, r, c)
        if v is None:
            continue
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            if float(v).is_integer() and 0 < v < 200:
                enum = enum or "num"
            continue
        if not isinstance(v, str) or not v.strip() or looks_like_period(v):
            continue
        k = enumerator_kind(v) if len(v.strip()) <= 5 else None
        if k:
            enum, token = k, v.strip().strip(".)(")
            continue
        label, label_col = v.strip(), c
    if label:
        inline = leading_enumerator(label)
        if inline and enumerator_kind(inline) and not enum:
            enum, token = enumerator_kind(inline), inline
        label = strip_leading_enumerators(label) or label
    if raw:
        return label, label_col, enum, token
    return label, label_col, enum


def _detect_bank_name(grid, sheet_name, first_period_col):
    """Nama bank biasanya ada di atas kolom periode pertama (mis. F5 'BANK CENTRAL ASIA')."""
    for r in range(3, 8):
        for c in range(first_period_col, first_period_col + 4):
            v = _cell(grid, r, c)
            if isinstance(v, str) and not looks_like_period(v) and _is_heading(v) and not v.strip()[0].isdigit():
                return v.strip()
    return sheet_name


def build_sheet_model(name, grid):
    header_row, hits = find_header_row(grid)
    if header_row is None or hits < 3:
        return None
    period_cols = {}
    for c in range(1, MAX_COLS + 1):
        v = _cell(grid, header_row, c)
        if looks_like_period(v):
            period_cols[c] = str(v).strip()
    if not period_cols:
        return None
    first_pc = min(period_cols)

    # ---- deteksi seksi --------------------------------------------------------
    heads = {}
    for r in range(1, len(grid) + 1):
        for c in range(1, first_pc):
            v = _cell(grid, r, c)
            if not isinstance(v, str) or not _is_heading(v):
                continue
            u = v.upper()
            for key, pat in (("CAP", _HEAD_CAP), ("COMMIT", _HEAD_COMMIT), ("IS", _HEAD_IS), ("BS", _HEAD_BS)):
                if key not in heads and pat.search(u):
                    heads[key] = r
                    break
    if "BS" not in heads:
        return None

    order = sorted(heads.items(), key=lambda kv: kv[1])
    sections = {}
    for i, (key, start) in enumerate(order):
        if key == "COMMIT":
            continue
        end = order[i + 1][1] - 1 if i + 1 < len(order) else None
        if key == "CAP":
            last = None
            for r in range(start, min(start + 130, len(grid)) + 1):
                lbl, _, _ = _row_texts(grid, r, first_pc)
                if lbl and _CAP_LAST_ROW.search(lbl):
                    last = r
            end = last if last else min(start + 72, len(grid))
        if end is None:
            end = min(start + 110, len(grid))
        sections[key] = (start, end)

    # ---- baris akun ---------------------------------------------------------------
    rows = []
    for key, (start, end) in sections.items():
        stack = []  # (level, label, row)
        last_letter = None
        for r in range(start + 1, end + 1):
            label, label_col, enum, token = _row_texts(grid, r, first_pc, raw=True)
            if not label or label.strip().lower() in _SKIP_LABELS:
                continue
            # 'i'/'v'/'x' bisa huruf urut ATAU angka romawi -> lihat huruf sebelumnya
            if enum == "letter" and token and token.lower() in ("i", "v", "x"):
                if last_letter != chr(ord(token.lower()) - 1):
                    enum = "roman"
            if enum == "letter" and token:
                last_letter = token.lower()
            # baris lanjutan (label terpotong ke baris berikut, huruf kecil, tanpa nomor)
            glued = re.match(r"^([a-z])([A-Z]\w+)", label)  # 'cInstrumen ...' -> enumerator nempel
            if glued:
                enum, token, label = "letter", glued.group(1), label[1:]
            if enum is None and label[:1].islower() and rows and rows[-1].section == key and rows[-1].row >= r - 1:
                prev = rows[-1]
                prev.label = f"{prev.label} {label}"
                for c in range(first_pc, MAX_COLS + 1):
                    v = _cell(grid, r, c)
                    if v is not None and prev.cells.get(c) is None:
                        prev.cells[c] = v
                continue
            level = {"num": 1, "letter": 2, "roman": 3}.get(enum, 1)
            if token and enum in ("letter", "roman") and token.isupper():
                level = 0  # 'A', 'B', 'II' = judul blok (mis. 'A Modal Inti')
            if _is_heading(label) and enum is None:
                level = 0
            while stack and stack[-1][0] >= level:
                stack.pop()
            parent, parent_row = (stack[-1][1], stack[-1][2]) if stack and level > 0 else (None, None)
            if level > 0:
                stack.append((level, label, r))
            cells = {c: _cell(grid, r, c) for c in range(first_pc, MAX_COLS + 1)}
            cells = {c: v for c, v in cells.items() if v is not None}
            rows.append(TemplateRow(row=r, label=label, section=key, level=level, parent=parent,
                                    parent_row=parent_row, cells=cells))

    # ---- baris spesial di luar seksi (dipakai saat menambah kolom baru) ------------
    special = {}
    cap_end = max(e for _, e in sections.values())
    for r in range(cap_end + 1, len(grid) + 1):
        label, _, _ = _row_texts(grid, r, first_pc)
        if label and label.strip().upper() == "ANNUALIZER":
            special["annualizer"] = r
        if "period_date" not in special:
            vals = [_cell(grid, r, c) for c in period_cols]
            if sum(1 for v in vals if hasattr(v, "year") and hasattr(v, "month")) >= 2:
                special["period_date"] = r

    return SheetModel(
        name=name,
        bank_name=_detect_bank_name(grid, name, first_pc),
        header_row=header_row,
        period_cols=period_cols,
        sections=sections,
        rows=rows,
        max_row=len(grid),
        special_rows=special,
    )


def load_bank_models(path_or_file):
    """Baca semua sheet bank dari template (mode read-only -> jauh lebih cepat
    daripada load penuh). Return dict nama_sheet -> SheetModel."""
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path_or_file, read_only=True, data_only=False)
    models = {}
    try:
        for ws in wb.worksheets:
            first = list(ws.iter_rows(min_row=1, max_row=6, max_col=10, values_only=True))
            if not is_bank_sheet(first):
                continue
            model = build_sheet_model(ws.title, grid_from_worksheet(ws))
            if model and {"BS", "IS", "CAP"} <= set(model.sections):
                models[ws.title] = model
    finally:
        wb.close()
    return models
