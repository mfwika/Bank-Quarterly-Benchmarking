"""Laporan publikasi format OJK (export Excel dari portal cfs.ojk.go.id).

Satu laporan bank per periode di-download sebagai beberapa file Excel terpisah (Posisi Keuangan,
Laba Rugi, KPMM, Kualitas Aset, Rasio). Formatnya sama untuk semua bank:

    A2  'Laporan Publikasi Triwulanan'
    A3  judul laporan
    A5  'Juni 2026'
    A7  'PT BANK MANDIRI (PERSERO) '
    ... header kolom (Individual/Konsolidasian x periode, atau L/DPK/KL/D/M/Jumlah)
    ... pos-pos di kolom A (bernomor, menjorok sesuai level), angka di kolom kanan

File-file milik bank & periode yang sama digabung jadi satu laporan (`merge_sources`).
"""
from __future__ import annotations

import io
import re
import warnings
from collections import defaultdict

from .periods import MONTHS, quarter_end
from .text import parse_number, strip_leading_enumerators

_TITLE_SECTIONS = [
    (re.compile(r"posisi keuangan|neraca", re.I), "BS"),
    (re.compile(r"laba rugi", re.I), "IS"),
    (re.compile(r"penyediaan modal|kpmm", re.I), "CAP"),
    (re.compile(r"komitmen", re.I), "COMMIT"),
]
_CATS = {"l": "L", "dpk": "DPK", "kl": "KL", "d": "D", "m": "M", "jumlah": "JUMLAH", "total": "JUMLAH"}
_ENTITY = {"individual": "IND", "konsolidasian": "KONS", "bank": "IND", "konsolidasi": "KONS"}


def _period(text):
    m = re.fullmatch(r"\s*([A-Za-z]+)\s+(\d{4})\s*", str(text or ""))
    if not m:
        return None
    mon = MONTHS.get(m.group(1).lower())
    return (mon, int(m.group(2))) if mon else None


def _load(file_bytes):
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)


def is_ojk_excel(file_bytes) -> bool:
    try:
        ws = _load(file_bytes).worksheets[0]
    except Exception:
        return False
    a2, a3 = str(ws.cell(2, 1).value or ""), str(ws.cell(3, 1).value or "")
    return bool(re.search(r"laporan publikasi", a2, re.I) and re.match(r"\s*laporan\b", a3, re.I)
                and _period(ws.cell(5, 1).value))


def _grid_with_merged(ws, max_row):
    """Nilai sel header dengan sel gabungan (merged) diisi ke semua kolomnya."""
    grid = {}
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                grid[(r, c)] = v
    for rng in ws.merged_cells.ranges:
        v = grid.get((rng.min_row, rng.min_col))
        if v is None or rng.min_row > max_row:
            continue
        for r in range(rng.min_row, min(rng.max_row, max_row) + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid.setdefault((r, c), v)
    return grid


def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return parse_number(str(v).strip())


def parse_ojk_excel(file_bytes: bytes):
    """-> (rows, text). Tiap SourceRow diberi `group` = 'OJK|<bank>|<bulan>/<tahun>' untuk digabung."""
    from .source_parser import SourceRow

    wb = _load(file_bytes)
    ws = wb.worksheets[0]
    title = str(ws.cell(3, 1).value or "").strip()
    period = _period(ws.cell(5, 1).value)
    bank = re.sub(r"\s+", " ", str(ws.cell(7, 1).value or "")).strip()
    section = next((s for pat, s in _TITLE_SECTIONS if pat.search(title)), "OTHER")
    is_aq = bool(re.search(r"kualitas aset", title, re.I))
    is_ratio = bool(re.search(r"rasio", title, re.I))

    # baris data pertama: label di kolom A dan ada angka di kanannya
    first = None
    for r in range(8, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() and any(
                _num(ws.cell(r, c).value) is not None for c in range(2, ws.max_column + 1)):
            first = r
            break
    if first is None:
        return [], "\n".join(x for x in (title, bank) if x)
    grid = _grid_with_merged(ws, first - 1)

    # header per kolom angka
    meta = {}
    for c in range(2, ws.max_column + 1):
        per, ent, cat = None, None, None
        for r in range(8, first):
            v = grid.get((r, c))
            if v is None:
                continue
            s = str(v).strip()
            per = per or _period(s)
            ent = ent or _ENTITY.get(s.lower())
            cat = cat or _CATS.get(s.lower())
        if per or cat:
            meta[c] = (per, ent or ("IND" if (is_aq or is_ratio) else None), cat)
    # kolom yang hanya 'ikut' header gabungan (mis. 'M' di J:K) dan tidak pernah berisi angka -> buang
    data_rows = range(first, ws.max_row + 1)
    vcols = [c for c in sorted(meta) if any(_num(ws.cell(r, c).value) is not None for r in data_rows)]
    if not vcols:
        return [], title

    group = f"OJK|{bank.upper()}|{period[0]}/{period[1]}" if period else None
    company = bank
    rows, stack = [], []  # stack: (indent, label)
    for r in range(first, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        if raw is None or not str(raw).strip():
            continue
        raw = str(raw).replace("\xa0", " ")
        indent = len(raw) - len(raw.lstrip(" "))
        label = strip_leading_enumerators(re.sub(r"^\s*([IVX]+|\d+)(\s*\.\s*\d+)*\s*\.?\s+", "", raw.strip())
                                          if re.match(r"^\s*\d+(\s*\.\s*\d+)+\s", raw) else raw.strip())
        label = re.sub(r"\s+", " ", label).strip(" .:")
        if not re.search(r"[A-Za-z]{2,}", label):
            continue
        while stack and stack[-1][0] >= indent:
            stack.pop()
        ancestors = [lbl for _, lbl in stack]
        vals = [_num(ws.cell(r, c).value) for c in vcols]
        stack.append((indent, label))
        if all(v is None for v in vals):
            continue  # judul kelompok (induk baris berikutnya)
        values = [0.0 if v is None else v for v in vals]
        cols = [(meta[c][0], meta[c][1]) for c in vcols]
        cats = [meta[c][2] for c in vcols] if is_aq else None
        if cats:
            values, cols, cats = _repair_categories(values, cols, cats)
        row = SourceRow(len(rows), label, values, section, ancestors[-1] if ancestors else None, 1, cols)
        row.ancestors = ancestors
        row.cats = cats
        row.company = company
        row.group = group
        rows.append(row)
    day = quarter_end(*period).day if period else 30
    mon_name = str(ws.cell(5, 1).value or "").strip()
    text = "\n".join([title, f"Laporan Publikasi per {day} {mon_name}", company] + [r.label for r in rows])
    return rows, text


def _repair_categories(values, cols, cats):
    """Export OJK kadang kehilangan satu kolom kolektibilitas per periode (Jun 26 tanpa D, Jun 25
    tanpa KL). Jumlah tetap ada -> kolom yang hilang = Jumlah - kolom lain."""
    full = ["L", "DPK", "KL", "D", "M"]
    by_per = defaultdict(list)
    for i, ((per, ent), cat) in enumerate(zip(cols, cats)):
        by_per[(per, ent)].append(i)
    values, cols, cats = list(values), list(cols), list(cats)
    for key, idxs in by_per.items():
        have = {cats[i] for i in idxs}
        missing = [c for c in full if c not in have]
        tot = next((values[i] for i in idxs if cats[i] == "JUMLAH"), None)
        if len(missing) == 1 and tot is not None:
            values.append(tot - sum(values[i] for i in idxs if cats[i] in full))
            cols.append(key)
            cats.append(missing[0])
    return values, cols, cats


def merge_sources(parts):
    """parts: list of (rows, text) -> (rows, text) dengan idx berurutan (urutan: BS, IS, CAP, lainnya)."""
    order = {"BS": 0, "IS": 1, "CAP": 2, "COMMIT": 3}
    parts = sorted(parts, key=lambda p: min((order.get(r.section, 9) for r in p[0]), default=9))
    rows, texts = [], []
    for prow, ptext in parts:
        for r in prow:
            r.idx = len(rows)
            rows.append(r)
        texts.append(ptext)
    return rows, "\n".join(texts)
