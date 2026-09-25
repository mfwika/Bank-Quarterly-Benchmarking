"""Tulis hasil ke workbook template: isi / tambah kolom periode."""
import copy
from datetime import datetime

from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

from .matcher import translate_formula
from .periods import quarter_end
from .template_model import SheetModel, is_ref_formula


def _copy_style(src, dst):
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.border = copy.copy(src.border)
        dst.fill = copy.copy(src.fill)
        dst.number_format = src.number_format
        dst.alignment = copy.copy(src.alignment)
        dst.protection = copy.copy(src.protection)


def apply_to_workbook(wb, model: SheetModel, target_col: int, period_label: str, month: int, year: int,
                      values: dict, notes: dict | None = None):
    """
    - Kolom baru (periode belum ada): salin format + perpanjang SEMUA formula
      dari kolom periode sebelumnya (termasuk rasio-rasio di bawah), lalu isi angka.
    - Kolom sudah ada (mis. 'Dec 25' sudah dibuat tapi kosong): formula yang
      belum ada dilengkapi, angka diisi.
    values: row -> angka. notes: row -> teks komentar (opsional, utk sel yang perlu dicek).
    Return jumlah sel angka yang ditulis.
    """
    ws = wb[model.name]
    ref_col = model.reference_col(target_col)
    is_new = target_col not in model.period_cols

    if is_new:
        src_letter, dst_letter = get_column_letter(ref_col), get_column_letter(target_col)
        if src_letter in ws.column_dimensions:
            ws.column_dimensions[dst_letter].width = ws.column_dimensions[src_letter].width

    header = ws.cell(model.header_row, target_col)
    if is_new or header.value in (None, ""):
        header.value = period_label
        _copy_style(ws.cell(model.header_row, ref_col), header)

    for r in range(1, ws.max_row + 1):
        if r == model.header_row:
            continue
        src = ws.cell(r, ref_col)
        dst = ws.cell(r, target_col)
        if is_new:
            _copy_style(src, dst)
        if dst.value is None and is_ref_formula(src.value):
            dst.value = translate_formula(src.value, ref_col, target_col, r)

    ann = model.special_rows.get("annualizer")
    if ann:
        ws.cell(ann, target_col).value = f"=12/{month}"
    date_row = model.special_rows.get("period_date")
    if date_row:
        d = quarter_end(month, year)
        ws.cell(date_row, target_col).value = datetime(d.year, d.month, d.day)

    written = 0
    for r, v in values.items():
        if v is None:
            continue
        cell = ws.cell(r, target_col)
        cell.value = v
        written += 1
        if notes and notes.get(r):
            cell.comment = Comment(notes[r], "Auto-Update Benchmarking")
    return written
