import io
import os
import warnings

import openpyxl
import pytest

from bankbench.aliases import merged_aliases
from bankbench.matcher import Matcher, reconcile
from bankbench.periods import detect_report_period, format_period_label, parse_period_label
from bankbench.source_parser import parse_source, split_label_values
from bankbench.template_model import build_sheet_model, grid_from_worksheet, load_bank_models
from bankbench.text import clean_label, parse_number, similarity
from bankbench.writer import apply_to_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "template.xlsx")


# ---------------------------------------------------------------- text / angka
@pytest.mark.parametrize("raw,expected", [
    ("1.234.567", 1234567), ("(1.234)", -1234), ("1,234,567", 1234567), ("12,50", 12.5),
    ("-", 0.0), ("-5.000", -5000), ("29,69%", 29.69), ("0", 0.0),
])
def test_parse_number(raw, expected):
    assert parse_number(raw) == pytest.approx(expected)


def test_split_label_values_keeps_numbers_inside_label():
    tok, label, vals = split_label_values("II. MODAL PELENGKAP (TIER 2) 9.604.882 (8.923) - 10.110.402")
    assert tok == "II"
    assert label == "MODAL PELENGKAP (TIER 2)"
    assert vals == [9604882, -8923, 0.0, 10110402]


def test_clean_label_and_similarity():
    assert clean_label("K a s") == "kas"
    assert clean_label("a Tahun-tahun lalu") == "tahun tahun lalu"
    # repo vs reverse repo harus dibedakan
    repo = "Surat berharga yang dijual dengan janji dibeli kembali (repo)"
    rev = "Tagihan atas surat berharga yang dibeli dengan janji dijual kembali (reverse repo)"
    assert similarity(repo, repo) > similarity(repo, rev) + 5


def test_periods():
    assert parse_period_label("Sep 25") == (9, 2025)
    assert parse_period_label("Mar 2009") == (3, 2009)
    assert format_period_label(12, 2025, like="Sep 25") == "Dec 25"
    assert detect_report_period("Per 30 September 2025 dan 31 Desember 2024") == (9, 2025)


# ---------------------------------------------------------------- mini template
def _mini_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BANKX"
    ws["B2"] = "NERACA - Consol"
    ws["F5"] = "BANK XYZ"
    for c, lbl in zip("FGH", ["Jun 25", "Sep 25", "Dec 25"]):
        ws[f"{c}6"] = lbl
    rows = [
        (8, "1", "K a s", 100_000, 110_000),
        (9, "2", "Kredit", None, None),
        (10, "a", "Pinjaman yang diberikan dan piutang", 500_000, 520_000),
        (11, "3", "Aset lainnya", 50_000, 55_000),
    ]
    for r, enum, lbl, v1, v2 in rows:
        ws[f"B{r}" if enum.isdigit() else f"D{r}"] = enum
        ws[f"D{r}" if enum.isdigit() else f"E{r}"] = lbl
        ws[f"F{r}"], ws[f"G{r}"] = v1, v2
    ws["E12"] = "TOTAL ASET"
    ws["F12"], ws["G12"] = "=SUM(F8:F11)", "=SUM(G8:G11)"
    ws["B20"] = "PERHITUNGAN LABA RUGI - Consol"
    ws["B22"], ws["C22"], ws["F22"], ws["G22"] = 1, "Pendapatan Bunga", 7_000, 11_000
    ws["B23"], ws["C23"], ws["F23"], ws["G23"] = 2, "Beban Bunga", 2_000, 3_000
    ws["C24"], ws["F24"], ws["G24"] = "Pendapatan Bunga Bersih", "=+F22-F23", "=+G22-G23"
    ws["B30"] = "KOMITMEN DAN KONTIGENSI - Consol"
    ws["B40"] = "CAPITAL ADEQUACY RATIO - Consol"
    ws["C42"], ws["F42"], ws["G42"] = "Aset Tertimbang Menurut Risiko (ATMR) untuk Risiko Kredit", 400_000, 410_000
    ws["C43"], ws["F43"], ws["G43"] = "Rasio Kewajiban Penyediaan Modal Minimum", "=F42/2", "=G42/2"
    ws["E50"], ws["F50"], ws["G50"] = "ANNUALIZER", "=12/6", "=12/9"
    return wb


def _mini_model(wb):
    return build_sheet_model("BANKX", grid_from_worksheet(wb["BANKX"]))


def test_template_model_sections_and_hierarchy():
    m = _mini_model(_mini_workbook())
    assert set(m.sections) == {"BS", "IS", "CAP"}
    assert m.bank_name == "BANK XYZ"
    rows = {r.row: r for r in m.rows}
    assert rows[10].parent == "Kredit"
    assert m.target_col_for(12, 2025) == (8, True)  # 'Dec 25' sudah ada (kolom H)
    assert m.target_col_for(3, 2026) == (9, False)  # kolom baru
    assert m.special_rows["annualizer"] == 50


def _source_rows(text):
    from bankbench.source_parser import _feed_text, _RowBuilder

    b = _RowBuilder()
    _feed_text(b, text)
    return b.rows


REPORT = """PT BANK X Tbk
LAPORAN POSISI KEUANGAN
Per 31 Desember 2025 dan 31 Desember 2024
1. Kas 90.000 80.000 120.000 95.000
2. Kredit yang diberikan 480.000 400.000 530.000 450.000
3. Aset lainnya 50.000 40.000 56.000 52.000
TOTAL ASET 620.000 520.000 706.000 597.000
LAPORAN LABA RUGI
1. Pendapatan Bunga 10.000 9.000 12.000 11.000
2. Beban Bunga (3.000) (2.500) (3.100) (3.000)
Pendapatan (Beban) Bunga Bersih 7.000 6.500 8.900 8.000
LAPORAN PERHITUNGAN KEWAJIBAN PENYEDIAAN MODAL MINIMUM
ATMR RISIKO KREDIT 390.000 380.000 420.000 410.000
"""


def test_matcher_identifies_accounts_and_column():
    m = _mini_model(_mini_workbook())
    src = _source_rows(REPORT)
    mt = Matcher(m, src, 8, merged_aliases({}))
    res = {r.row: r for r in mt.run()}
    assert mt.layout.cur_idx == 2  # Konsolidasian periode ini
    assert res[8].value == 120_000
    assert res[10].value == 530_000  # 'Kredit yang diberikan' -> sub-baris 'Pinjaman ...'
    assert res[9].value is None  # induk 'Kredit' tidak diisi (hindari dobel hitung)
    assert res[11].value == 56_000
    assert res[23].value == 3_100  # laporan (3.100) -> template simpan beban positif
    assert res[42].value == 420_000
    assert res[12].kind == "Formula"
    checks = {c["Baris"]: c for c in reconcile(mt, {r: x.value for r, x in res.items() if x.value is not None})}
    assert checks[12]["Status"] == "OK"


def test_writer_new_column_extends_formulas():
    wb = _mini_workbook()
    m = _mini_model(wb)
    n = apply_to_workbook(wb, m, 9, "Mar 26", 3, 2026, {8: 1.0, 10: 2.0}, {8: "cek"})
    ws = wb["BANKX"]
    assert n == 2
    assert ws["I6"].value == "Mar 26"
    assert ws["I12"].value == "=SUM(I8:I11)"
    assert ws["I24"].value == "=+I22-I23"
    assert ws["I50"].value == "=12/3"
    assert ws["I8"].comment.text == "cek"


# ---------------------------------------------------------------- end-to-end (template asli)
@pytest.mark.slow
@pytest.mark.skipif(not os.path.exists(TEMPLATE), reason="template.xlsx tidak ada")
@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
def test_end_to_end_real_template(tmp_path, fmt):
    pytest.importorskip("reportlab")
    from tests.sample_report import build_report_lines, write_pdf, write_xlsx

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        models = load_bank_models(TEMPLATE)
    m = models["BCA"]
    blocks, expected = build_report_lines(m, 19, 16, 15)  # Sep 25 vs Dec 24 / Sep 24
    path = tmp_path / f"bca.{fmt}"
    (write_pdf if fmt == "pdf" else write_xlsx)(blocks, path)
    rows, text = parse_source(str(path), path.read_bytes())
    mt = Matcher(m, rows, 19, merged_aliases({}))
    res = {r.row: r for r in mt.run()}
    wrong = {r: (e, res[r].value) for r, e in expected.items() if res[r].value is None or abs(res[r].value - e) > 1}
    assert not wrong
    checks = reconcile(mt, {r: x.value for r, x in res.items() if x.value is not None})
    balance = [c for c in checks if c["Total (template)"].startswith("Aset =")]
    assert balance and balance[0]["Status"] == "OK"
