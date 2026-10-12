import io
from datetime import datetime

import openpyxl

from bankbench.compare import STATUS_DIFF, STATUS_ONLY_ANALYST, compare_workbooks
from bankbench.writer import ISSUE_SHEET, write_issue_sheet


def _wb(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bank"
    ws.cell(1, 2).value, ws.cell(1, 3).value = "Dec 25", "Jun 26"
    for r, (label, prev, cur) in enumerate(rows, start=2):
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value = label, prev, cur
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_compare_flags_diff_missing_and_stale_copy():
    ours = _wb([("Kas", 1, 100), ("Giro", 1, 50), ("Total", None, "=SUM(C2:C3)"), ("NPL", "=(1+2)/1", "=(3+4)/1"),
                ("LCR", 1.2, None)])
    analyst = _wb([("Kas", 1, 100), ("Giro", 1, 55), ("Total", None, 155), ("NPL", "=(1+2)/1", "=(1+2)/1"),
                   ("LCR", 1.2, 1.3)])
    rows, summary = compare_workbooks(ours, analyst, [dict(sheet="Bank", col=3, label="Jun 26", header_row=1)])
    by = {r["Akun"]: r for r in rows}
    assert "Kas" not in by
    assert by["Giro"]["Status"] == STATUS_DIFF and by["Giro"]["Selisih"] == -5
    assert by["Total"]["Status"] == STATUS_DIFF  # 150 (=SUM) vs 155 diketik analis
    assert "tersalin" in by["NPL"]["Keterangan"]
    assert by["LCR"]["Status"] == STATUS_ONLY_ANALYST
    assert summary[0]["Sama"] == 1


def test_compare_missing_sheet_or_column():
    ours = _wb([("Kas", 1, 100)])
    rows, summary = compare_workbooks(ours, ours, [dict(sheet="Lain", col=3, label="Jun 26")])
    assert not rows and "tidak ada" in summary[0]["Keterangan"]
    rows, summary = compare_workbooks(ours, ours, [dict(sheet="Bank", col=3, label="Sep 26")])
    assert "tidak ditemukan" in summary[0]["Keterangan"]


def test_issue_sheet_written():
    wb = openpyxl.Workbook()
    write_issue_sheet(wb, [{"File": "a.pdf", "Sheet": "BCA", "Periode": "Jun 26", "Baris": 46, "Akun": "TOTAL ASET",
                            "Masalah": "BEDA dengan laporan", "Keterangan": "x"}], datetime(2026, 10, 12))
    ws = wb[ISSUE_SHEET]
    assert ws.cell(3, 5).value == "TOTAL ASET" and ws.cell(3, 6).value == "BEDA dengan laporan"
    write_issue_sheet(wb, [], datetime(2026, 10, 12))  # ditulis ulang, tidak dobel
    assert wb.sheetnames.count(ISSUE_SHEET) == 1
