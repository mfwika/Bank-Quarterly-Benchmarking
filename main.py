"""Auto-Update Template Benchmarking Bank.

Alur:
1. Template (template.xlsx) sudah tersimpan di repo / backend.
2. User upload laporan keuangan publikasi (PDF / Excel) — boleh beberapa bank sekaligus.
3. Sistem otomatis: tebak bank (sheet), tebak periode, baca akun Neraca, Laba Rugi &
   Struktur Modal (KPMM), identifikasi akun ke baris template (alias + pola angka +
   nama akun), pilih kolom angka yang benar (konsolidasian / periode berjalan).
4. User review & koreksi di tabel (opsional) -> cek rekonsiliasi total.
5. Generate: kolom periode di-isi / ditambah (formula ikut diperpanjang) -> download.
"""
import hashlib
import io
import os
import warnings
from datetime import datetime

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from bankbench.aliases import MAPPINGS_PATH, add_learned, load_learned, merged_aliases, save_learned
from bankbench.banks import detect_bank, sheet_display
from bankbench.matcher import Matcher, reconcile
from bankbench.periods import detect_report_period, format_period_label
from bankbench.source_parser import parse_source
from bankbench.template_model import SECTION_NAMES, load_bank_models
from bankbench.writer import apply_to_workbook

st.set_page_config(page_title="Auto-Update Template Benchmarking", layout="wide")
st.title("Auto-Update Template Benchmarking")
st.caption(
    "Upload laporan keuangan publikasi bank → sistem mengenali bank, periode & akun "
    "(Neraca, Laba Rugi, Struktur Modal) → review → template ter-update dengan kolom periode baru."
)

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "template.xlsx")
if not os.path.exists(TEMPLATE_PATH):
    st.error(
        "File template belum ada. Taruh file Excel template dengan nama persis **template.xlsx** "
        "di folder yang sama dengan main.py (di repo), lalu commit."
    )
    st.stop()
TEMPLATE_MTIME = os.path.getmtime(TEMPLATE_PATH)

QUARTER_MONTHS = [3, 6, 9, 12]
QUARTER_NAMES = {3: "Q1 (Mar)", 6: "Q2 (Jun)", 9: "Q3 (Sep)", 12: "Q4 (Dec)"}
SOURCE_NONE = "(kosong / tidak diisi)"


# ---------------------------------------------------------------------------------
# cache
# ---------------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def get_models(path, mtime):
    return load_bank_models(path)


@st.cache_data(show_spinner=False)
def get_template_bytes(path, mtime):
    with open(path, "rb") as fh:
        return fh.read()


@st.cache_data(show_spinner=False, max_entries=30)
def parse_cached(name, data):
    return parse_source(name, data)


def file_key(f):
    return hashlib.md5(f.getvalue()).hexdigest()[:10]


def fmt_num(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return f"{v:,.0f}".replace(",", ".")


with st.spinner("Membaca struktur template (sekali saja, lalu di-cache)..."):
    models = get_models(TEMPLATE_PATH, TEMPLATE_MTIME)
if not models:
    st.error("Tidak ada sheet bank (berisi NERACA / LABA RUGI / CAPITAL ADEQUACY) yang terdeteksi di template.")
    st.stop()

if "learned" not in st.session_state:
    st.session_state.learned = load_learned()
aliases = merged_aliases(st.session_state.learned)

with st.sidebar:
    st.header("ℹ️ Cara kerja")
    st.markdown(
        "- **Seksi yang diisi**: Neraca, Laba Rugi, Struktur Modal (KPMM/CAR). "
        "Komitmen & kontinjensi dan rasio-rasio (formula) tidak disentuh, formulanya ikut diperpanjang.\n"
        "- **Identifikasi akun**:\n"
        "  1. *Pola angka* — angka pembanding di laporan (mis. Des tahun lalu) dicocokkan dengan "
        "histori di template → akun & kolom (konsolidasian/bank only) ketahuan otomatis.\n"
        "  2. *Alias* — kamus padanan nama akun (bisa belajar dari koreksimu).\n"
        "  3. *Nama akun* — kemiripan nama + induk akun + urutan.\n"
        "- Baris yang **biasanya kosong** di template tidak diisi otomatis (hindari dobel hitung)."
    )
    st.divider()
    st.caption(f"Template: `template.xlsx` — {len(models)} sheet bank terdeteksi.")
    n_learned = sum(len(v) for v in st.session_state.learned.values())
    st.caption(f"Pemetaan hasil belajar: {n_learned} akun (`mappings.json`).")

# ---------------------------------------------------------------------------------
# STEP 1: upload
# ---------------------------------------------------------------------------------
uploads = st.file_uploader(
    "① Drag & drop laporan keuangan publikasi (PDF / Excel) — boleh beberapa bank sekaligus",
    type=["pdf", "xlsx", "xlsm"],
    accept_multiple_files=True,
)
if not uploads:
    st.info("Upload minimal satu laporan publikasi untuk mulai.")
    st.stop()

parsed = {}
for f in uploads:
    with st.spinner(f"Membaca '{f.name}'..."):
        try:
            rows, text = parse_cached(f.name, f.getvalue())
        except Exception as e:  # file rusak / terenkripsi / bukan PDF teks
            st.error(f"Gagal membaca '{f.name}': {e}")
            continue
    relevant = [r for r in rows if r.section in ("BS", "IS", "CAP", "?")]
    if not relevant:
        st.error(
            f"Tidak ada akun + angka yang terbaca dari '{f.name}'. Kalau PDF-nya hasil scan (gambar), "
            "sistem tidak bisa membaca teksnya — pakai versi PDF teks / Excel dari website bank / OJK."
        )
        continue
    parsed[file_key(f)] = (f, rows, text)
if not parsed:
    st.stop()

# ---------------------------------------------------------------------------------
# STEP 2: periode
# ---------------------------------------------------------------------------------
detected = [detect_report_period(text) for _, _, text in parsed.values()]
detected = [d for d in detected if d]
any_model = next(iter(models.values()))
if detected:
    default_period = max(set(detected), key=detected.count)
else:
    latest = max((m.latest_period() for m in models.values() if m.latest_period()), key=lambda p: (p[1], p[0]))
    mon, yr = latest
    default_period = (mon + 3, yr) if mon < 12 else (3, yr + 1)

st.subheader("② Periode laporan")
c1, c2, c3 = st.columns([1, 1, 2])
with c1:
    month = st.selectbox("Kuartal", QUARTER_MONTHS, index=QUARTER_MONTHS.index(default_period[0])
                         if default_period[0] in QUARTER_MONTHS else 3, format_func=lambda m: QUARTER_NAMES[m])
with c2:
    year = st.number_input("Tahun", min_value=2000, max_value=2100, value=int(default_period[1]), step=1)
period_label = format_period_label(month, int(year), like=any_model.period_cols[any_model.last_period_col()])
with c3:
    st.markdown(
        f"Judul kolom: **{period_label}**"
        + (f"  \n<small>Terdeteksi dari isi laporan: {', '.join(f'{m:02d}/{y}' for m, y in sorted(set(detected)))}</small>"
           if detected else "  \n<small>Periode tidak terbaca dari laporan — pilih manual.</small>"),
        unsafe_allow_html=True,
    )

# ---------------------------------------------------------------------------------
# STEP 3: per file -> bank, matching, review
# ---------------------------------------------------------------------------------
st.subheader("③ Review hasil identifikasi akun")
jobs = {}
sheet_names = list(models)
tabs = st.tabs([f.name for f, _, _ in parsed.values()])


def build_df(matcher, results, rows_by_idx, option_of):
    recs = []
    for r in results:
        src_opt = option_of.get(r.source_idx, SOURCE_NONE) if r.source_idx is not None else SOURCE_NONE
        recs.append({
            "Baris": r.row,
            "Seksi": r.section,
            "Akun (Template)": ("   ↳ " if r.parent else "") + r.label,
            "Tipe": r.kind,
            "Akun di Laporan": src_opt,
            "Nilai": r.value,
            "Periode lalu": r.prev_value,
            "Skor": r.score if r.source_idx is not None else None,
            "Metode": r.method,
            "Catatan": r.note,
        })
    df = pd.DataFrame(recs).set_index("Baris", drop=False)
    return df


def on_editor_change(state_key, editor_key, shown_index, matcher_key):
    """Terapkan editan ke tabel utama. Kalau 'Akun di Laporan' diganti, nilainya ikut
    diambil otomatis dari akun laporan yang dipilih + dicatat sebagai pemetaan baru."""
    df = st.session_state[state_key]
    ctx = st.session_state[matcher_key]
    edits = st.session_state.get(editor_key, {}).get("edited_rows", {})
    for pos, changes in edits.items():
        row_id = shown_index[int(pos)]
        for col, val in changes.items():
            if col == "Akun di Laporan":
                df.at[row_id, col] = val
                sr = ctx["src_by_option"].get(val)
                tr = ctx["rows_by_num"][row_id]
                if sr is None:
                    df.at[row_id, "Nilai"] = None
                else:
                    v, notes = ctx["matcher"].value_for(tr, sr)
                    df.at[row_id, "Nilai"] = v
                    add_learned(st.session_state.learned, tr.section, tr.label, tr.parent, sr.label)
                df.at[row_id, "Metode"] = "Manual"
                df.at[row_id, "Skor"] = None
            elif col == "Nilai":
                df.at[row_id, "Nilai"] = val
                df.at[row_id, "Metode"] = "Manual"
    st.session_state[state_key] = df
    st.session_state.pop(editor_key, None)


for tab, (fk, (f, rows, text)) in zip(tabs, parsed.items()):
    with tab:
        ranking = detect_bank(text, rows, models, aliases)
        best_sheet = ranking[0][0]
        reason = {n: r for n, _, r in ranking}
        cc1, cc2 = st.columns([2, 3])
        with cc1:
            sheet = st.selectbox(
                "Bank / sheet tujuan",
                sheet_names,
                index=sheet_names.index(best_sheet),
                format_func=lambda n: sheet_display(models[n]),
                key=f"sheet_{fk}",
            )
        with cc2:
            st.caption(f"Tebakan otomatis: **{best_sheet}** — {reason[best_sheet]}")
        model = models[sheet]
        target_col, exists = model.target_col_for(month, int(year))
        ref_col = model.reference_col(target_col)
        col_letter = get_column_letter(target_col)
        if exists and model.fill_ratio(target_col) > 0.05:
            st.warning(f"Kolom **{period_label}** (kolom {col_letter}) di sheet ini SUDAH BERISI data → angka lama "
                       "akan ditimpa dengan angka dari laporan ini. Pastikan periodenya benar.")
        elif exists:
            st.info(f"Kolom **{period_label}** sudah ada di sheet ini (kolom {col_letter}) → akan diisi. "
                    f"Pola/format mengikuti kolom {get_column_letter(ref_col)} ('{model.period_cols.get(ref_col)}').")
        else:
            st.success(f"Kolom baru **{period_label}** akan ditambahkan di kolom {col_letter} "
                       f"(setelah '{model.period_cols[model.last_period_col()]}'); formula & format disalin "
                       f"dari kolom {get_column_letter(ref_col)}.")

        matcher = Matcher(model, rows, target_col, aliases)
        layout = matcher.detect_layout()
        with st.expander("🔢 Kolom angka yang dipakai dari laporan", expanded=layout.comp_idx is None):
            if layout.comp_idx is not None:
                comp_txt = ", ".join(f"{SECTION_NAMES.get(s, s)} → '{p}'" for s, p in layout.comp_period_by_section.items())
                st.markdown(
                    f"Tiap baris laporan punya **{layout.n_values} angka**. Angka ke-**{layout.comp_idx + 1}** "
                    f"persis sama dengan histori template ({comp_txt}) → itu kolom **pembanding**. "
                    f"Jadi angka ke-**{layout.cur_idx + 1}** dipakai sebagai **periode berjalan**."
                    + (f" Skala dikonversi ×{layout.scale:g}." if layout.scale != 1 else "")
                )
            else:
                st.warning(
                    "Pola angka pembanding tidak ditemukan di histori template (bank baru / format beda). "
                    f"Sistem menebak angka ke-{layout.cur_idx + 1} sebagai periode berjalan"
                    + (" (kolom Konsolidasian periode ini di format OJK)." if layout.n_values >= 4 else ".")
                    + " Cek & ubah di bawah kalau salah."
                )
            cur_idx = st.selectbox(
                "Angka ke berapa (dari kiri) = periode berjalan?",
                list(range(max(layout.n_values, 1))),
                index=min(layout.cur_idx, max(layout.n_values, 1) - 1),
                format_func=lambda i: f"Angka ke-{i + 1}",
                key=f"cur_{fk}",
            )
            preview = [
                {"Seksi": r.section, "Akun (laporan)": r.label, **{f"#{i + 1}": fmt_num(v) for i, v in enumerate(r.values[:6])}}
                for r in matcher.source[:400]
            ]
            st.dataframe(pd.DataFrame(preview), width="stretch", height=220)

        results = matcher.run(cur_idx_override=cur_idx)
        rows_by_idx = {r.idx: r for r in matcher.source}
        option_of = {r.idx: f"#{r.idx} · {r.label[:70]} [{fmt_num(matcher.current_value(r))}]" for r in matcher.source}
        src_by_option = {opt: rows_by_idx[i] for i, opt in option_of.items()}

        sig = (fk, sheet, target_col, cur_idx, int(year), month)
        state_key, matcher_key = f"df_{fk}", f"matcher_{fk}"
        if st.session_state.get(f"sig_{fk}") != sig:
            st.session_state[state_key] = build_df(matcher, results, rows_by_idx, option_of)
            st.session_state[f"sig_{fk}"] = sig
            st.session_state.pop(f"editor_{fk}", None)
        st.session_state[matcher_key] = {
            "matcher": matcher,
            "src_by_option": src_by_option,
            "rows_by_num": {r.row: r for r in model.rows},
        }
        df = st.session_state[state_key]

        n_val = int((df["Tipe"] == "Nilai").sum())
        n_filled = int(((df["Tipe"] != "Formula") & df["Nilai"].notna()).sum())
        n_check = int(((df["Tipe"] == "Nilai") & (df["Nilai"].isna() | df["Catatan"].str.contains("cek", na=False))).sum())
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Akun yang biasa diisi", n_val)
        m2.metric("Terisi otomatis", n_filled)
        m3.metric("Perlu dicek", n_check)
        m4.metric("Cocok via pola angka", int((df["Metode"] == "Pola angka").sum()))

        fc1, fc2 = st.columns([3, 2])
        with fc1:
            secs = st.multiselect("Seksi", ["BS", "IS", "CAP"], default=["BS", "IS", "CAP"],
                                  format_func=lambda s: SECTION_NAMES[s], key=f"secs_{fk}")
        with fc2:
            view = st.radio("Tampilkan", ["Akun yang diisi", "Perlu dicek saja", "Semua baris"],
                            horizontal=True, key=f"view_{fk}")
        shown = df[df["Seksi"].isin(secs)]
        if view == "Akun yang diisi":
            shown = shown[(shown["Tipe"] == "Nilai") | shown["Nilai"].notna()]
        elif view == "Perlu dicek saja":
            shown = shown[(shown["Tipe"] == "Nilai") & (shown["Nilai"].isna() | shown["Catatan"].str.contains("cek", na=False))]

        st.data_editor(
            shown,
            key=f"editor_{fk}",
            on_change=on_editor_change,
            args=(state_key, f"editor_{fk}", list(shown.index), matcher_key),
            width="stretch",
            height=460,
            hide_index=True,
            disabled=["Baris", "Seksi", "Akun (Template)", "Tipe", "Periode lalu", "Skor", "Metode", "Catatan"],
            column_config={
                "Baris": st.column_config.NumberColumn(width="small"),
                "Akun di Laporan": st.column_config.SelectboxColumn(
                    options=[SOURCE_NONE] + list(option_of.values()), width="large"),
                "Nilai": st.column_config.NumberColumn(format="%.0f"),
                "Periode lalu": st.column_config.NumberColumn(format="%.0f"),
                "Skor": st.column_config.ProgressColumn(min_value=0, max_value=100, format="%.0f"),
            },
        )
        st.caption("Ganti **Akun di Laporan** → nilai ikut terisi otomatis & pemetaannya diingat. "
                   "Atau ketik langsung di kolom **Nilai**. Baris 'Formula' dihitung otomatis oleh Excel.")

        values = {int(r): (None if pd.isna(v) else float(v)) for r, v in df["Nilai"].items()
                  if df.at[r, "Tipe"] != "Formula"}
        values = {r: v for r, v in values.items() if v is not None}
        checks = reconcile(matcher, values)
        if checks:
            chk = pd.DataFrame(checks)
            bad = int((chk["Status"] == "BEDA").sum())
            with st.expander(f"✅ Cek total vs laporan — {len(chk) - bad} cocok, {bad} beda", expanded=bad > 0):
                show = chk.copy()
                for c in ("Hasil hitung", "Di laporan", "Selisih"):
                    show[c] = show[c].map(fmt_num)
                st.dataframe(show, width="stretch", hide_index=True)
                st.caption("Total dihitung dari formula template memakai angka yang akan diisi, lalu dibandingkan "
                           "dengan total di laporan. 'BEDA' = ada akun yang salah pasang / belum terisi.")

        notes = {}
        for r, row in df.iterrows():
            if r in values and (row["Metode"] == "Manual" or (row["Skor"] is not None and not pd.isna(row["Skor"])
                                                              and row["Skor"] < 90 and row["Metode"] == "Nama akun")):
                notes[int(r)] = f"Auto-update: {row['Metode']} — {row['Akun di Laporan']}"
        jobs[fk] = dict(file=f.name, sheet=sheet, model=model, target_col=target_col, values=values, notes=notes)

# ---------------------------------------------------------------------------------
# STEP 4: generate
# ---------------------------------------------------------------------------------
st.subheader("④ Generate template")
dupes = pd.Series([j["sheet"] for j in jobs.values()]).value_counts()
dupes = dupes[dupes > 1]
if len(dupes):
    st.warning(f"Beberapa file diarahkan ke sheet yang sama: {', '.join(dupes.index)} — file terakhir yang menang.")

annotate = st.checkbox("Beri comment di sel yang diisi manual / skornya rendah (biar gampang dicek di Excel)", value=True)
if st.button(f"🚀 Update template — kolom '{period_label}' untuk {len(jobs)} bank", type="primary"):
    with st.spinner("Memuat template lengkap & menulis kolom baru... (±30–60 detik untuk file besar)"):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(get_template_bytes(TEMPLATE_PATH, TEMPLATE_MTIME)))
        summary = []
        for job in jobs.values():
            n = apply_to_workbook(wb, job["model"], job["target_col"], period_label, month, int(year),
                                  job["values"], job["notes"] if annotate else None)
            summary.append({"File": job["file"], "Sheet": job["sheet"],
                            "Kolom": get_column_letter(job["target_col"]), "Sel angka terisi": n})
        out = io.BytesIO()
        wb.save(out)
    st.session_state["output"] = (out.getvalue(), f"template_{period_label.replace(' ', '')}_{datetime.now():%Y%m%d}.xlsx",
                                  summary)

if "output" in st.session_state:
    data, fname, summary = st.session_state["output"]
    st.dataframe(pd.DataFrame(summary), hide_index=True, width="stretch")
    st.download_button("⬇️ Download template hasil update (.xlsx)", data=data, file_name=fname,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.caption("Catatan: openpyxl tidak menyimpan ulang chart/gambar yang ada di workbook. "
               "Kalau template punya chart, salin sheet hasil update ke template aslimu.")

if st.session_state.learned:
    with st.expander("🧠 Pemetaan akun yang dipelajari dari koreksimu"):
        st.json(st.session_state.learned, expanded=False)
        b1, b2 = st.columns(2)
        with b1:
            if st.button("💾 Simpan ke mappings.json (server)"):
                save_learned(st.session_state.learned)
                st.success(f"Tersimpan di {MAPPINGS_PATH}. Di Streamlit Cloud file ini hilang saat app restart — "
                           "download lalu commit ke repo supaya permanen.")
        with b2:
            import json

            st.download_button("⬇️ Download mappings.json", data=json.dumps(st.session_state.learned, indent=2,
                                                                           ensure_ascii=False),
                               file_name="mappings.json", mime="application/json")
