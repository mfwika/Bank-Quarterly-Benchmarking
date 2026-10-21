"""Build the practice Excel templates for IB Prep Studio.

    python ib-prep/build_templates.py          # write ib-prep/templates/*.xlsx
    python ib-prep/build_templates.py --embed  # embed those files into ib-prep/index.html

All numbers are illustrative (fictional "ExampleCo"). Every output cell is a live formula,
so changing a blue input recalculates the whole model.
"""
import base64
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "templates")

# ---------- styling ----------
F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10)
GREEN = Font(name=F, size=10, color="008000")
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True, color="A3122A")
SUB = Font(name=F, size=10, italic=True, color="666666")
HDR = Font(name=F, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="A3122A")
YEAR_FILL = PatternFill("solid", fgColor="F2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
BASE_FILL = PatternFill("solid", fgColor="FCE4E8")
TOP = Border(top=Side(style="thin", color="000000"))

NUM = '#,##0.0_);(#,##0.0);"-"_)'
INT = '#,##0_);(#,##0);"-"_)'
USD = '$#,##0.00_);($#,##0.00);"-"_)'
PCT = '0.0%_);(0.0%);"-"_)'
MULT = '0.0"x"_);(0.0"x");"-"_)'
BETA = '0.00_);(0.00);"-"_)'

SRC = "Source: illustrative figures for a fictional company (ExampleCo). Replace with real filings (10-K / annual report)."


def put(ws, ref, value, kind="f", fmt=None, bold=False, fill=None):
    """kind: 'in' = hard-coded input (blue), 'key' = key input (blue on yellow),
    'f' = formula (black), 'link' = link to another sheet (green), 'lbl' = label."""
    c = ws[ref]
    c.value = value
    font = {"in": BLUE, "key": BLUE, "f": BLACK, "link": GREEN, "lbl": BLACK}[kind]
    if bold:
        font = Font(name=F, size=10, bold=True, color=font.color)
    c.font = font
    if kind == "key":
        c.fill = INPUT_FILL
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def title(ws, text, sub):
    ws["A1"].value = text
    ws["A1"].font = TITLE
    ws["A2"].value = sub
    ws["A2"].font = SUB
    ws["A3"].value = "Legend: blue = input, yellow fill = key assumption to change, black = formula, green = link to another sheet."
    ws["A3"].font = SUB
    ws.sheet_view.showGridLines = False


def section(ws, row, text, last_col):
    for col in range(1, last_col + 1):
        c = ws.cell(row=row, column=col)
        c.fill = HDR_FILL
        c.font = HDR
    ws.cell(row=row, column=1).value = text


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def guide(wb, name, purpose, steps, drills):
    ws = wb.active
    ws.title = "Guide"
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 4, "B": 110})
    ws["B1"].value = name
    ws["B1"].font = TITLE
    ws["B2"].value = "IB Prep Studio practice template. " + SRC
    ws["B2"].font = SUB
    r = 4
    blocks = [("What this model does", purpose), ("How to use it", steps), ("Practice drills", drills),
              ("Colour legend", [
                  "Blue text = hard-coded input. Yellow fill = key assumption you are meant to change.",
                  "Black text = formula. Never type over a black cell.",
                  "Green text = link pulling a value from another sheet.",
                  "Negative numbers show in (parentheses); zeros show as '-'.",
              ])]
    for head, lines in blocks:
        ws.cell(row=r, column=2, value=head).font = Font(name=F, size=11, bold=True, color="A3122A")
        r += 1
        for i, line in enumerate(lines, 1):
            c = ws.cell(row=r, column=2, value=f"{i}. {line}")
            c.font = BLACK
            c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1
    # example of each style
    for i, (lbl, kind) in enumerate([("Input example", "in"), ("Key assumption example", "key"),
                                     ("Formula example", "f"), ("Link example", "link")]):
        put(ws, f"B{r + i}", lbl, kind)


# =====================================================================
# 1. Three-statement model
# =====================================================================
def three_statement():
    wb = Workbook()
    guide(wb, "Three-Statement Model",
          ["Projects the income statement, balance sheet and cash flow statement for 5 years from 2 years of history.",
           "Links the statements through working capital, PP&E and debt schedules, with a revolver that draws when cash falls below the minimum.",
           "Interest is calculated on beginning-of-year balances, so the model has no circular references."],
          ["Go to the 'Model' sheet. Historical columns (FY2024A-FY2025A) are blue hard-codes; projections (FY2026E-FY2030E) are formulas.",
           "Change the yellow assumptions (growth, margins, days, capex, debt repayment, payout, minimum cash).",
           "Watch the balance check row: it must read OK in every year. If it doesn't, something is mis-linked.",
           "Best practice: rebuild this model on a blank sheet without looking, then compare it with this one."],
          ["Set FY2026E revenue growth to 20%. Before looking, predict whether cash goes up or down, and why (hint: working capital).",
           "Raise AR days from 50 to 70. Which three lines move on the cash flow statement and balance sheet?",
           "Set the minimum cash to 200. Find the year the revolver draws and trace where the draw shows up.",
           "Set the dividend payout to 100%. What happens to retained earnings and to the revolver?",
           "Delete the formula in one projected BS cell, see the check break, then fix it."])
    ws = wb.create_sheet("Model")
    title(ws, "ExampleCo - Three-Statement Model", "($ in millions, fiscal years ending 31-Dec)   " + SRC)
    widths(ws, {"A": 40, **{c: 12 for c in "BCDEFGH"}})
    hist, proj = ["B", "C"], ["D", "E", "F", "G", "H"]
    years = ["FY2024A", "FY2025A", "FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E"]
    allc = hist + proj
    prev = lambda c: chr(ord(c) - 1)

    layout = [
        ("#", "ASSUMPTIONS"),
        ("growth", "Revenue growth"), ("gm", "Gross margin"), ("sga", "SG&A % of revenue"),
        ("da", "D&A % of revenue"), ("capex", "Capex % of revenue"), ("tax", "Tax rate"),
        ("ar_days", "Accounts receivable days (of revenue)"), ("inv_days", "Inventory days (of COGS)"),
        ("ap_days", "Accounts payable days (of COGS)"), ("int_debt", "Interest rate - term debt"),
        ("int_rev", "Interest rate - revolver"), ("int_cash", "Interest rate earned on cash"),
        ("mand", "Mandatory term debt repayment ($)"), ("payout", "Dividend payout (% of net income)"),
        ("min_cash", "Minimum cash balance ($)"),
        ("", ""),
        ("#", "INCOME STATEMENT"),
        ("rev", "Revenue"), ("cogs", "Cost of goods sold"), ("gp", "Gross profit"), ("sga_x", "SG&A"),
        ("ebitda", "EBITDA"), ("da_x", "Depreciation & amortization"), ("ebit", "EBIT"),
        ("int_x", "Interest expense"), ("int_i", "Interest income"), ("ebt", "Pre-tax income"),
        ("taxes", "Income taxes"), ("ni", "Net income"), ("ebitda_m", "EBITDA margin"),
        ("", ""),
        ("#", "BALANCE SHEET"),
        ("cash", "Cash"), ("ar", "Accounts receivable"), ("inv", "Inventory"), ("tca", "Total current assets"),
        ("ppe", "PP&E, net"), ("ta", "Total assets"), ("", ""),
        ("ap", "Accounts payable"), ("revolver", "Revolver"), ("debt", "Term debt"), ("tl", "Total liabilities"),
        ("common", "Common stock & APIC"), ("re", "Retained earnings"), ("te", "Total equity"),
        ("tle", "Total liabilities & equity"), ("check", "Balance check (assets - L&E)"),
        ("", ""),
        ("#", "CASH FLOW STATEMENT"),
        ("cf_ni", "Net income"), ("cf_da", "Plus: D&A"), ("cf_ar", "(Increase) / decrease in receivables"),
        ("cf_inv", "(Increase) / decrease in inventory"), ("cf_ap", "Increase / (decrease) in payables"),
        ("cfo", "Cash flow from operations"), ("cf_capex", "Capital expenditures"), ("cfi", "Cash flow from investing"),
        ("cf_repay", "Term debt repayment"), ("cf_div", "Dividends paid"), ("cf_rev", "Revolver draw / (repayment)"),
        ("cff", "Cash flow from financing"), ("net", "Net change in cash"), ("beg_cash", "Beginning cash"),
        ("end_cash", "Ending cash"),
        ("", ""),
        ("#", "SUPPORTING SCHEDULES"),
        ("ppe_beg", "PP&E - beginning"), ("ppe_capex", "Plus: capex"), ("ppe_da", "Less: D&A"),
        ("ppe_end", "PP&E - ending"), ("", ""),
        ("debt_beg", "Term debt - beginning"), ("debt_rep", "Less: mandatory repayment"), ("debt_end", "Term debt - ending"),
        ("", ""),
        ("cash_pre", "Cash before revolver (beginning cash + CFO + CFI + repayment + dividends)"),
        ("rev_beg", "Revolver - beginning"), ("rev_draw", "Draw / (repayment)"), ("rev_end", "Revolver - ending"),
    ]
    R, row = {}, 6
    for key, _ in layout:
        if key and key != "#":
            R[key] = row
        row += 1

    # headers
    ws["A5"].value = "($ in millions)"
    ws["A5"].font = BOLD
    for c, y in zip(allc, years):
        cell = ws[f"{c}5"]
        cell.value = y
        cell.font = BOLD
        cell.fill = YEAR_FILL
        cell.alignment = Alignment(horizontal="right")
    row = 6
    for key, label in layout:
        if key == "#":
            section(ws, row, label, 8)
        elif key:
            ws.cell(row=row, column=1, value=label).font = BLACK
        row += 1

    hist_data = {
        "rev": [1000, 1100], "cogs": [600, 655], "sga_x": [200, 215], "da_x": [40, 44],
        "int_x": [24, 22], "int_i": [2, 3], "taxes": [34.5, 41.75],
        "cash": [80, 95], "ar": [137, 151], "inv": [99, 108], "ppe": [400, 420],
        "ap": [74, 81], "revolver": [0, 0], "debt": [300, 280], "common": [150, 150], "re": [192, 263],
    }
    for k, vals in hist_data.items():
        for c, v in zip(hist, vals):
            put(ws, f"{c}{R[k]}", v, "in", NUM)

    assumptions = {
        "growth": ([0.08, 0.07, 0.06, 0.05, 0.04], PCT), "gm": ([0.405] * 5, PCT), "sga": ([0.195] * 5, PCT),
        "da": ([0.04] * 5, PCT), "capex": ([0.05] * 5, PCT), "tax": ([0.25] * 5, PCT),
        "ar_days": ([50] * 5, INT), "inv_days": ([60] * 5, INT), "ap_days": ([45] * 5, INT),
        "int_debt": ([0.075] * 5, PCT), "int_rev": ([0.08] * 5, PCT), "int_cash": ([0.03] * 5, PCT),
        "mand": ([20] * 5, NUM), "payout": ([0.30] * 5, PCT), "min_cash": ([50] * 5, NUM),
    }
    for k, (vals, fmt) in assumptions.items():
        for c, v in zip(proj, vals):
            put(ws, f"{c}{R[k]}", v, "key", fmt)
    # historical ratios (formulas, for reference)
    for c in hist:
        r = R
        if c == "C":
            put(ws, f"C{r['growth']}", f"=C{r['rev']}/B{r['rev']}-1", "f", PCT)
        put(ws, f"{c}{r['gm']}", f"={c}{r['gp']}/{c}{r['rev']}", "f", PCT)
        put(ws, f"{c}{r['sga']}", f"={c}{r['sga_x']}/{c}{r['rev']}", "f", PCT)
        put(ws, f"{c}{r['da']}", f"={c}{r['da_x']}/{c}{r['rev']}", "f", PCT)
        put(ws, f"{c}{r['tax']}", f"={c}{r['taxes']}/{c}{r['ebt']}", "f", PCT)
        put(ws, f"{c}{r['ar_days']}", f"={c}{r['ar']}/{c}{r['rev']}*365", "f", INT)
        put(ws, f"{c}{r['inv_days']}", f"={c}{r['inv']}/{c}{r['cogs']}*365", "f", INT)
        put(ws, f"{c}{r['ap_days']}", f"={c}{r['ap']}/{c}{r['cogs']}*365", "f", INT)

    def f(key, fn, cols, fmt=NUM, bold=False, border=False):
        for c in cols:
            cell = put(ws, f"{c}{R[key]}", fn(c, prev(c)), "f", fmt, bold)
            if border:
                cell.border = TOP
        if bold:
            ws[f"A{R[key]}"].font = BOLD

    r = R
    # income statement
    f("rev", lambda c, p: f"={p}{r['rev']}*(1+{c}{r['growth']})", proj, bold=True)
    f("cogs", lambda c, p: f"={c}{r['rev']}*(1-{c}{r['gm']})", proj)
    f("gp", lambda c, p: f"={c}{r['rev']}-{c}{r['cogs']}", allc, border=True)
    f("sga_x", lambda c, p: f"={c}{r['rev']}*{c}{r['sga']}", proj)
    f("ebitda", lambda c, p: f"={c}{r['gp']}-{c}{r['sga_x']}", allc, bold=True, border=True)
    f("da_x", lambda c, p: f"={c}{r['ppe_da']}", proj)
    f("ebit", lambda c, p: f"={c}{r['ebitda']}-{c}{r['da_x']}", allc, border=True)
    f("int_x", lambda c, p: f"={c}{r['debt_beg']}*{c}{r['int_debt']}+{c}{r['rev_beg']}*{c}{r['int_rev']}", proj)
    f("int_i", lambda c, p: f"={p}{r['cash']}*{c}{r['int_cash']}", proj)
    f("ebt", lambda c, p: f"={c}{r['ebit']}-{c}{r['int_x']}+{c}{r['int_i']}", allc, border=True)
    f("taxes", lambda c, p: f"={c}{r['ebt']}*{c}{r['tax']}", proj)
    f("ni", lambda c, p: f"={c}{r['ebt']}-{c}{r['taxes']}", allc, bold=True, border=True)
    f("ebitda_m", lambda c, p: f"={c}{r['ebitda']}/{c}{r['rev']}", allc, PCT)
    # balance sheet
    f("cash", lambda c, p: f"={c}{r['end_cash']}", proj)
    f("ar", lambda c, p: f"={c}{r['rev']}*{c}{r['ar_days']}/365", proj)
    f("inv", lambda c, p: f"={c}{r['cogs']}*{c}{r['inv_days']}/365", proj)
    f("tca", lambda c, p: f"=SUM({c}{r['cash']}:{c}{r['inv']})", allc, border=True)
    f("ppe", lambda c, p: f"={c}{r['ppe_end']}", proj)
    f("ta", lambda c, p: f"={c}{r['tca']}+{c}{r['ppe']}", allc, bold=True, border=True)
    f("ap", lambda c, p: f"={c}{r['cogs']}*{c}{r['ap_days']}/365", proj)
    f("revolver", lambda c, p: f"={c}{r['rev_end']}", proj)
    f("debt", lambda c, p: f"={c}{r['debt_end']}", proj)
    f("tl", lambda c, p: f"=SUM({c}{r['ap']}:{c}{r['debt']})", allc, border=True)
    f("common", lambda c, p: f"={p}{r['common']}", proj)
    f("re", lambda c, p: f"={p}{r['re']}+{c}{r['ni']}+{c}{r['cf_div']}", proj)
    f("te", lambda c, p: f"={c}{r['common']}+{c}{r['re']}", allc, border=True)
    f("tle", lambda c, p: f"={c}{r['tl']}+{c}{r['te']}", allc, bold=True, border=True)
    f("check", lambda c, p: f'=IF(ABS({c}{r["ta"]}-{c}{r["tle"]})<0.001,"OK","ERROR")', allc, bold=True)
    for c in allc:
        ws[f"{c}{r['check']}"].alignment = Alignment(horizontal="right")
    # cash flow statement
    f("cf_ni", lambda c, p: f"={c}{r['ni']}", proj)
    f("cf_da", lambda c, p: f"={c}{r['da_x']}", proj)
    f("cf_ar", lambda c, p: f"={p}{r['ar']}-{c}{r['ar']}", proj)
    f("cf_inv", lambda c, p: f"={p}{r['inv']}-{c}{r['inv']}", proj)
    f("cf_ap", lambda c, p: f"={c}{r['ap']}-{p}{r['ap']}", proj)
    f("cfo", lambda c, p: f"=SUM({c}{r['cf_ni']}:{c}{r['cf_ap']})", proj, bold=True, border=True)
    f("cf_capex", lambda c, p: f"=-{c}{r['ppe_capex']}", proj)
    f("cfi", lambda c, p: f"={c}{r['cf_capex']}", proj, bold=True, border=True)
    f("cf_repay", lambda c, p: f"=-{c}{r['debt_rep']}", proj)
    f("cf_div", lambda c, p: f"=-MAX(0,{c}{r['ni']})*{c}{r['payout']}", proj)
    f("cf_rev", lambda c, p: f"={c}{r['rev_draw']}", proj)
    f("cff", lambda c, p: f"=SUM({c}{r['cf_repay']}:{c}{r['cf_rev']})", proj, bold=True, border=True)
    f("net", lambda c, p: f"={c}{r['cfo']}+{c}{r['cfi']}+{c}{r['cff']}", proj, bold=True)
    f("beg_cash", lambda c, p: f"={p}{r['cash']}", proj)
    f("end_cash", lambda c, p: f"={c}{r['beg_cash']}+{c}{r['net']}", proj, bold=True, border=True)
    # schedules
    f("ppe_beg", lambda c, p: f"={p}{r['ppe']}", proj)
    f("ppe_capex", lambda c, p: f"={c}{r['rev']}*{c}{r['capex']}", proj)
    f("ppe_da", lambda c, p: f"={c}{r['rev']}*{c}{r['da']}", proj)
    f("ppe_end", lambda c, p: f"={c}{r['ppe_beg']}+{c}{r['ppe_capex']}-{c}{r['ppe_da']}", proj, border=True)
    f("debt_beg", lambda c, p: f"={p}{r['debt']}", proj)
    f("debt_rep", lambda c, p: f"=MIN({c}{r['debt_beg']},{c}{r['mand']})", proj)
    f("debt_end", lambda c, p: f"={c}{r['debt_beg']}-{c}{r['debt_rep']}", proj, border=True)
    f("cash_pre", lambda c, p: f"={c}{r['beg_cash']}+{c}{r['cfo']}+{c}{r['cfi']}+{c}{r['cf_repay']}+{c}{r['cf_div']}", proj)
    f("rev_beg", lambda c, p: f"={p}{r['revolver']}", proj)
    f("rev_draw", lambda c, p: f"=MAX({c}{r['min_cash']}-{c}{r['cash_pre']},-{c}{r['rev_beg']})", proj)
    f("rev_end", lambda c, p: f"={c}{r['rev_beg']}+{c}{r['rev_draw']}", proj, border=True)
    ws.freeze_panes = "B6"
    return wb


# =====================================================================
# 2. DCF
# =====================================================================
def dcf():
    wb = Workbook()
    guide(wb, "DCF Model",
          ["Values ExampleCo with an unlevered DCF: 5-year UFCF projection, WACC built from CAPM and a peer beta table, and terminal value by both the perpetuity growth and exit multiple methods.",
           "Bridges enterprise value to equity value and an implied share price, and cross-checks each terminal value method against the other.",
           "Two sensitivity tables (WACC vs growth, WACC vs exit multiple) are built with plain formulas, so they update instantly."],
          ["Go to the 'DCF' sheet. Change the yellow inputs: market data, CAPM inputs, peer betas, projection drivers.",
           "Read the valuation block: compare the two methods and check the implied exit multiple and implied growth rate.",
           "Check 'PV of terminal value as % of EV'. If it is above ~80%, your projections carry little weight.",
           "Rebuild it from a blank sheet: drivers -> UFCF -> discount factors -> terminal value -> EV -> equity value -> per share."],
          ["Raise the equity risk premium by 1%. Which moves more: the perpetuity value or the exit multiple value? Why?",
           "Turn the mid-year convention off (set it to 0). Does value go up or down?",
           "Set terminal growth equal to WACC. What breaks, and why is that the right answer?",
           "Change peer D/E ratios and watch the relevered beta and WACC change.",
           "Make the implied perpetual growth from the exit multiple method land at about 2.5% by adjusting the multiple."])
    ws = wb.create_sheet("DCF")
    title(ws, "ExampleCo - Discounted Cash Flow Analysis", "($ in millions, except per share)   " + SRC)
    widths(ws, {"A": 42, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 4, "I": 14, "J": 13, "K": 13, "L": 12, "M": 15})

    section(ws, 5, "INPUTS", 3)
    inputs = [
        (6, "Current share price ($)", 25.0, USD), (7, "Diluted shares outstanding (mm)", 100, NUM),
        (8, "Total debt", 280, NUM), (9, "Cash & equivalents", 95, NUM), (10, "Tax rate", 0.25, PCT),
        (11, "Risk-free rate", 0.0425, PCT), (12, "Equity risk premium", 0.055, PCT),
        (14, "Pre-tax cost of debt", 0.075, PCT), (15, "Terminal (perpetual) growth rate", 0.025, PCT),
        (16, "Exit EV / EBITDA multiple", 9.0, MULT), (17, "Mid-year convention (1 = on, 0 = off)", 1, INT),
    ]
    for r, lbl, v, fmt in inputs:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", v, "key", fmt)
    ws["A13"].value = "Beta (relevered, from peer table)"
    put(ws, "B13", "=K13", "f", BETA)

    section(ws, 19, "WACC", 3)
    wacc = [
        (20, "Cost of equity (CAPM: Rf + beta x ERP)", "=B11+B13*B12", PCT),
        (21, "After-tax cost of debt", "=B14*(1-B10)", PCT),
        (22, "Market value of equity", "=B6*B7", NUM), (23, "Debt", "=B8", NUM),
        (24, "Equity / (debt + equity)", "=B22/(B22+B23)", PCT), (25, "Debt / (debt + equity)", "=B23/(B22+B23)", PCT),
        (26, "WACC", "=B24*B20+B25*B21", PCT),
    ]
    for r, lbl, fx, fmt in wacc:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fx, "f", fmt, bold=(r == 26))
    ws["A26"].font = BOLD

    # peer beta table
    for col, h in zip("IJKLM", ["Peer", "Levered beta", "Debt / equity", "Tax rate", "Unlevered beta"]):
        c = ws[f"{col}5"]
        c.value = h
        c.font = HDR
        c.fill = HDR_FILL
    peers = [("Peer A", 1.10, 0.30, 0.25), ("Peer B", 1.25, 0.50, 0.25), ("Peer C", 0.95, 0.20, 0.25), ("Peer D", 1.30, 0.60, 0.25)]
    for i, (n, b, de, t) in enumerate(peers):
        r = 6 + i
        put(ws, f"I{r}", n, "in")
        put(ws, f"J{r}", b, "key", BETA)
        put(ws, f"K{r}", de, "key", PCT)
        put(ws, f"L{r}", t, "key", PCT)
        put(ws, f"M{r}", f"=J{r}/(1+(1-L{r})*K{r})", "f", BETA)
    ws["I11"].value = "Median unlevered beta"
    put(ws, "K11", "=MEDIAN(M6:M9)", "f", BETA)
    ws["I12"].value = "ExampleCo debt / equity"
    put(ws, "K12", "=B8/B22", "f", PCT)
    ws["I13"].value = "Relevered beta"
    ws["I13"].font = BOLD
    put(ws, "K13", "=K11*(1+(1-B10)*K12)", "f", BETA, bold=True)
    ws["I14"].value = "Formula: unlevered = levered / (1 + (1 - t) x D/E)"
    ws["I14"].font = SUB

    # projections
    cols, yrs = ["B", "C", "D", "E", "F", "G"], ["FY2025A", "FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E"]
    proj = cols[1:]
    section(ws, 28, "PROJECTIONS", 7)
    for c, y in zip(cols, yrs):
        cell = ws[f"{c}29"]
        cell.value = y
        cell.font = BOLD
        cell.fill = YEAR_FILL
        cell.alignment = Alignment(horizontal="right")
    drivers = [(30, "Revenue growth", [0.08, 0.07, 0.06, 0.05, 0.04]), (31, "EBITDA margin", [0.215, 0.22, 0.225, 0.23, 0.23]),
               (32, "D&A % of revenue", [0.04] * 5), (33, "Capex % of revenue", [0.05] * 5),
               (34, "Net working capital % of change in revenue", [0.16] * 5)]
    for r, lbl, vals in drivers:
        ws[f"A{r}"].value = lbl
        for c, v in zip(proj, vals):
            put(ws, f"{c}{r}", v, "key", PCT)
    ws["A36"].value = "Revenue"
    ws["A37"].value = "EBITDA"
    ws["A38"].value = "Less: D&A"
    ws["A39"].value = "EBIT"
    ws["A40"].value = "NOPAT (EBIT x (1 - tax))"
    ws["A41"].value = "Plus: D&A"
    ws["A42"].value = "Less: capex"
    ws["A43"].value = "Less: increase in net working capital"
    ws["A44"].value = "Unlevered free cash flow"
    ws["A44"].font = BOLD
    ws["A46"].value = "Year"
    ws["A47"].value = "Discount period"
    ws["A48"].value = "Discount factor"
    ws["A49"].value = "PV of unlevered FCF"
    put(ws, "B36", 1100, "in", NUM)
    put(ws, "B37", 230, "in", NUM)
    put(ws, "B38", 44, "in", NUM)
    put(ws, "B39", "=B37-B38", "f", NUM)
    for c in proj:
        p = chr(ord(c) - 1)
        put(ws, f"{c}36", f"={p}36*(1+{c}30)", "f", NUM, bold=True)
        put(ws, f"{c}37", f"={c}36*{c}31", "f", NUM)
        put(ws, f"{c}38", f"={c}36*{c}32", "f", NUM)
        put(ws, f"{c}39", f"={c}37-{c}38", "f", NUM).border = TOP
        put(ws, f"{c}40", f"={c}39*(1-$B$10)", "f", NUM)
        put(ws, f"{c}41", f"={c}38", "f", NUM)
        put(ws, f"{c}42", f"=-{c}36*{c}33", "f", NUM)
        put(ws, f"{c}43", f"=-({c}36-{p}36)*{c}34", "f", NUM)
        put(ws, f"{c}44", f"=SUM({c}40:{c}43)", "f", NUM, bold=True).border = TOP
        put(ws, f"{c}46", 1 if c == "C" else f"={p}46+1", "in" if c == "C" else "f", INT)
        put(ws, f"{c}47", f"={c}46-0.5*$B$17", "f", '0.0')
        put(ws, f"{c}48", f"=1/(1+$B$26)^{c}47", "f", '0.000')
        put(ws, f"{c}49", f"={c}44*{c}48", "f", NUM)

    section(ws, 51, "VALUATION", 3)
    ws["B52"].value, ws["C52"].value = "Perpetuity growth", "Exit multiple"
    for c in "BC":
        ws[f"{c}52"].font = BOLD
        ws[f"{c}52"].alignment = Alignment(horizontal="right", wrap_text=True)
    rows = [
        (53, "Sum of PV of unlevered FCF", "=SUM($C$49:$G$49)", "=SUM($C$49:$G$49)", NUM),
        (54, "Terminal value (in FY2030E)", "=G44*(1+B15)/(B26-B15)", "=G37*B16", NUM),
        (55, "PV of terminal value", "=B54/(1+$B$26)^$G$46", "=C54/(1+$B$26)^$G$46", NUM),
        (56, "Enterprise value", "=B53+B55", "=C53+C55", NUM),
        (57, "Less: debt", "=-$B$8", "=-$B$8", NUM), (58, "Plus: cash", "=$B$9", "=$B$9", NUM),
        (59, "Equity value", "=SUM(B56:B58)", "=SUM(C56:C58)", NUM),
        (60, "Implied share price ($)", "=B59/$B$7", "=C59/$B$7", USD),
        (61, "Premium / (discount) to current price", "=B60/$B$6-1", "=C60/$B$6-1", PCT),
        (62, "PV of terminal value as % of EV", "=B55/B56", "=C55/C56", PCT),
    ]
    for r, lbl, fb, fc, fmt in rows:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fb, "f", fmt, bold=r in (56, 60))
        put(ws, f"C{r}", fc, "f", fmt, bold=r in (56, 60))
    ws["A56"].font = BOLD
    ws["A60"].font = BOLD
    ws["A63"].value = "Cross-check: implied exit multiple / implied perpetual growth"
    put(ws, "B63", "=B54/G37", "f", MULT)
    put(ws, "C63", "=(C54*B26-G44)/(C54+G44)", "f", PCT)
    ws["A64"].value = "Note: terminal value is discounted from the end of FY2030E in both methods."
    ws["A64"].font = SUB

    # sensitivities
    deltas = [-0.01, -0.005, 0, 0.005, 0.01]
    section(ws, 66, "SENSITIVITY: SHARE PRICE ($), PERPETUITY GROWTH METHOD  (rows = WACC, columns = terminal growth)", 7)
    for j, d in enumerate(deltas):
        col = "BCDEF"[j]
        put(ws, f"{col}67", f"=$B$15+({d})", "f", PCT, bold=True).fill = YEAR_FILL
        put(ws, f"A{68 + j}", f"=$B$26+({d})", "f", PCT, bold=True).fill = YEAR_FILL
    for i in range(5):
        r = 68 + i
        for j in range(5):
            col = "BCDEF"[j]
            fx = (f"=(SUMPRODUCT($C$44:$G$44,1/(1+$A{r})^$C$47:$G$47)"
                  f"+$G$44*(1+{col}$67)/($A{r}-{col}$67)/(1+$A{r})^$G$46-$B$8+$B$9)/$B$7")
            put(ws, f"{col}{r}", fx, "f", USD, fill=BASE_FILL if (i == 2 and j == 2) else None)
    section(ws, 74, "SENSITIVITY: SHARE PRICE ($), EXIT MULTIPLE METHOD  (rows = WACC, columns = exit EV/EBITDA)", 7)
    for j, d in enumerate([-2, -1, 0, 1, 2]):
        col = "BCDEF"[j]
        put(ws, f"{col}75", f"=$B$16+({d})", "f", MULT, bold=True).fill = YEAR_FILL
        put(ws, f"A{76 + j}", f"=$B$26+({deltas[j]})", "f", PCT, bold=True).fill = YEAR_FILL
    for i in range(5):
        r = 76 + i
        for j in range(5):
            col = "BCDEF"[j]
            fx = (f"=(SUMPRODUCT($C$44:$G$44,1/(1+$A{r})^$C$47:$G$47)"
                  f"+$G$37*{col}$75/(1+$A{r})^$G$46-$B$8+$B$9)/$B$7")
            put(ws, f"{col}{r}", fx, "f", USD, fill=BASE_FILL if (i == 2 and j == 2) else None)
    ws["A81"].value = "Shaded cell = base case; it should equal the implied share price above."
    ws["A81"].font = SUB
    for rr in list(range(68, 73)) + list(range(76, 81)):
        ws[f"A{rr}"].alignment = Alignment(horizontal="right")
    return wb


# =====================================================================
# 3. LBO
# =====================================================================
def lbo():
    wb = Workbook()
    guide(wb, "LBO Model",
          ["A 5-year leveraged buyout of ExampleCo: sources & uses, operating projections, a debt schedule with mandatory amortization and a cash sweep, and sponsor returns (MOIC and IRR).",
           "Includes a returns attribution (EBITDA growth vs multiple expansion vs debt paydown) that ties to the equity gain, and an IRR sensitivity on entry vs exit multiples.",
           "Interest is charged on beginning-of-year balances, so there is no circularity."],
          ["Go to the 'LBO' sheet. Change the yellow inputs: entry multiple, leverage, rates, operating drivers, exit multiple.",
           "Sources & uses: sponsor equity is the plug that makes sources equal uses.",
           "Watch the term loan get swept down with excess cash, and net debt / EBITDA fall each year.",
           "Rebuild it on a blank sheet, then try a paper LBO with the same inputs and compare your mental-math IRR to the model's."],
          ["Raise leverage from 5.0x to 6.0x total. What happens to MOIC, IRR and year-1 interest coverage?",
           "Cut the exit multiple by 1.0x. How many IRR points do you lose?",
           "Set the cash sweep to 0%. Where does the excess cash go, and does the IRR change? (Hint: net debt.)",
           "Make the deal hit a 20% IRR by changing only the entry multiple.",
           "Look at the attribution: which driver creates most of the value in the base case?"])
    ws = wb.create_sheet("LBO")
    title(ws, "ExampleCo - Leveraged Buyout Analysis", "($ in millions)   " + SRC)
    widths(ws, {"A": 46, **{c: 13 for c in "BCDEFG"}, "H": 4, "I": 28, "J": 13, "K": 14})

    section(ws, 5, "TRANSACTION INPUTS", 2)
    inputs = [
        (6, "LTM revenue", 1100, NUM, "in"), (7, "LTM EBITDA", 230, NUM, "in"),
        (8, "Entry EV / EBITDA", 9.0, MULT, "key"), (9, "Transaction fees (% of EV)", 0.02, PCT, "key"),
        (10, "Cash to balance sheet at close", 0, NUM, "key"),
        (11, "Term loan B (x EBITDA)", 3.0, MULT, "key"), (12, "Term loan B interest rate", 0.08, PCT, "key"),
        (13, "Term loan B amortization (% of original per year)", 0.01, PCT, "key"),
        (14, "Cash sweep (% of excess cash used to prepay TLB)", 1.0, PCT, "key"),
        (15, "Senior notes (x EBITDA)", 2.0, MULT, "key"), (16, "Senior notes interest rate", 0.095, PCT, "key"),
        (17, "Tax rate", 0.25, PCT, "key"), (18, "Exit EV / EBITDA (end of year 5)", 9.0, MULT, "key"),
    ]
    for r, lbl, v, fmt, kind in inputs:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", v, kind, fmt)

    for col in "IJK":
        ws[f"{col}5"].fill = HDR_FILL
        ws[f"{col}5"].font = HDR
    ws["I5"].value, ws["J5"].value, ws["K5"].value = "Sources & uses", "$mm", "% / x EBITDA"
    su = [
        (6, "Uses", None, None), (7, "Purchase enterprise value", "=B7*B8", "=J7/$B$7"),
        (8, "Transaction fees", "=J7*B9", None), (9, "Cash to balance sheet", "=B10", None),
        (10, "Total uses", "=SUM(J7:J9)", None), (12, "Sources", None, None),
        (13, "Term loan B", "=B7*B11", "=J13/$J$16"), (14, "Senior notes", "=B7*B15", "=J14/$J$16"),
        (15, "Sponsor equity (plug)", "=J10-J13-J14", "=J15/$J$16"), (16, "Total sources", "=SUM(J13:J15)", "=J16/$J$16"),
    ]
    for r, lbl, fj, fk in su:
        ws[f"I{r}"].value = lbl
        if fj is None:
            ws[f"I{r}"].font = BOLD
            continue
        put(ws, f"J{r}", fj, "f", NUM, bold=r in (10, 16, 15))
        if fk:
            put(ws, f"K{r}", fk, "f", MULT if r == 7 else PCT)
    ws["I10"].font = BOLD
    ws["I16"].font = BOLD
    ws["I17"].value = "Check: sources - uses"
    put(ws, "J17", '=IF(ABS(J16-J10)<0.001,"OK","ERROR")', "f")
    ws["J17"].alignment = Alignment(horizontal="right")

    cols = ["B", "C", "D", "E", "F", "G"]
    proj = cols[1:]
    section(ws, 20, "OPERATING MODEL", 7)
    for c, y in zip(cols, ["Close (LTM)", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]):
        cell = ws[f"{c}21"]
        cell.value = y
        cell.font = BOLD
        cell.fill = YEAR_FILL
        cell.alignment = Alignment(horizontal="right")
    drivers = [(22, "Revenue growth", [0.07, 0.07, 0.06, 0.06, 0.05]), (23, "EBITDA margin", [0.21, 0.215, 0.22, 0.225, 0.23]),
               (24, "D&A % of revenue", [0.04] * 5), (25, "Capex % of revenue", [0.045] * 5),
               (26, "Net working capital % of change in revenue", [0.15] * 5)]
    for r, lbl, vals in drivers:
        ws[f"A{r}"].value = lbl
        for c, v in zip(proj, vals):
            put(ws, f"{c}{r}", v, "key", PCT)
    labels = {28: "Revenue", 29: "EBITDA", 30: "Less: D&A", 31: "EBIT", 32: "Less: interest - term loan B",
              33: "Less: interest - senior notes", 34: "Pre-tax income", 35: "Less: taxes", 36: "Net income",
              38: "Plus: D&A", 39: "Less: capex", 40: "Less: increase in net working capital",
              41: "Free cash flow available for debt repayment"}
    for r, lbl in labels.items():
        ws[f"A{r}"].value = lbl
    ws["A41"].font = BOLD
    put(ws, "B28", "=B6", "f", NUM)
    put(ws, "B29", "=B7", "f", NUM)
    for c in proj:
        p = chr(ord(c) - 1)
        put(ws, f"{c}28", f"={p}28*(1+{c}22)", "f", NUM, bold=True)
        put(ws, f"{c}29", f"={c}28*{c}23", "f", NUM, bold=True)
        put(ws, f"{c}30", f"={c}28*{c}24", "f", NUM)
        put(ws, f"{c}31", f"={c}29-{c}30", "f", NUM).border = TOP
        put(ws, f"{c}32", f"={c}45*$B$12", "f", NUM)
        put(ws, f"{c}33", f"={c}50*$B$16", "f", NUM)
        put(ws, f"{c}34", f"={c}31-{c}32-{c}33", "f", NUM).border = TOP
        put(ws, f"{c}35", f"=MAX(0,{c}34)*$B$17", "f", NUM)
        put(ws, f"{c}36", f"={c}34-{c}35", "f", NUM, bold=True).border = TOP
        put(ws, f"{c}38", f"={c}30", "f", NUM)
        put(ws, f"{c}39", f"={c}28*{c}25", "f", NUM)
        put(ws, f"{c}40", f"=({c}28-{p}28)*{c}26", "f", NUM)
        put(ws, f"{c}41", f"={c}36+{c}38-{c}39-{c}40", "f", NUM, bold=True).border = TOP

    section(ws, 43, "DEBT SCHEDULE", 7)
    dl = {45: "Term loan B - beginning", 46: "Mandatory amortization", 47: "Optional prepayment (cash sweep)",
          48: "Term loan B - ending", 50: "Senior notes - beginning", 51: "Senior notes - ending (bullet, repaid at exit)",
          53: "Cash - beginning", 54: "Change in cash", 55: "Cash - ending",
          57: "Total debt", 58: "Net debt", 59: "Net debt / EBITDA", 60: "EBITDA / cash interest"}
    for r, lbl in dl.items():
        ws[f"A{r}"].value = lbl
    put(ws, "B48", "=J13", "f", NUM)
    put(ws, "B51", "=J14", "f", NUM)
    put(ws, "B55", "=J9", "f", NUM)
    put(ws, "B57", "=B48+B51", "f", NUM)
    put(ws, "B58", "=B57-B55", "f", NUM)
    put(ws, "B59", "=B58/B29", "f", MULT)
    for c in proj:
        p = chr(ord(c) - 1)
        put(ws, f"{c}45", f"={p}48", "f", NUM)
        put(ws, f"{c}46", f"=-MIN({c}45,$J$13*$B$13)", "f", NUM)
        put(ws, f"{c}47", f"=-MIN({c}45+{c}46,MAX(0,{c}53+{c}41+{c}46)*$B$14)", "f", NUM)
        put(ws, f"{c}48", f"=SUM({c}45:{c}47)", "f", NUM, bold=True).border = TOP
        put(ws, f"{c}50", f"={p}51", "f", NUM)
        put(ws, f"{c}51", f"={c}50", "f", NUM, bold=True)
        put(ws, f"{c}53", f"={p}55", "f", NUM)
        put(ws, f"{c}54", f"={c}41+{c}46+{c}47", "f", NUM)
        put(ws, f"{c}55", f"={c}53+{c}54", "f", NUM, bold=True).border = TOP
        put(ws, f"{c}57", f"={c}48+{c}51", "f", NUM)
        put(ws, f"{c}58", f"={c}57-{c}55", "f", NUM)
        put(ws, f"{c}59", f"={c}58/{c}29", "f", MULT)
        put(ws, f"{c}60", f"={c}29/({c}32+{c}33)", "f", MULT)
    ws["A61"].value = "Note: cash earns no interest and there is no revolver; if FCF turns negative, cash goes negative."
    ws["A61"].font = SUB

    section(ws, 63, "RETURNS", 2)
    ret = [
        (64, "Exit-year EBITDA", "=G29", NUM), (65, "Exit enterprise value", "=B64*B18", NUM),
        (66, "Less: net debt at exit", "=-G58", NUM), (67, "Exit equity value", "=B65+B66", NUM),
        (68, "Sponsor equity invested", "=J15", NUM), (69, "MOIC", "=B67/B68", MULT),
        (70, "IRR (MOIC ^ (1/5) - 1)", "=B69^(1/5)-1", PCT), (71, "IRR (Excel IRR on cash flows below)", "=IRR(B73:G73)", PCT),
    ]
    for r, lbl, fx, fmt in ret:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fx, "f", fmt, bold=r in (67, 69, 70))
    ws["A73"].value = "Sponsor cash flows (Close -> Year 5)"
    put(ws, "B73", "=-B68", "f", NUM)
    for c in "CDEF":
        put(ws, f"{c}73", 0, "in", NUM)
    put(ws, "G73", "=B67", "f", NUM)

    section(ws, 75, "RETURNS ATTRIBUTION (where the equity gain comes from)", 2)
    attr = [
        (76, "EBITDA growth x entry multiple", "=(G29-B29)*B8"), (77, "Multiple expansion x exit EBITDA", "=(B18-B8)*G29"),
        (78, "Debt paydown & cash build", "=B58-G58"), (79, "Less: transaction fees", "=-J8"),
        (80, "Total equity value created", "=SUM(B76:B79)"),
    ]
    for r, lbl, fx in attr:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fx, "f", NUM, bold=(r == 80))
        if r < 80:
            put(ws, f"C{r}", f"=B{r}/$B$80", "f", PCT)
    ws["B80"].border = TOP
    ws["A81"].value = "Check: equals exit equity - equity invested"
    put(ws, "B81", '=IF(ABS(B80-(B67-B68))<0.001,"OK","ERROR")', "f")
    ws["B81"].alignment = Alignment(horizontal="right")

    section(ws, 83, "SENSITIVITY: IRR  (rows = entry EV/EBITDA, columns = exit EV/EBITDA)", 7)
    d = [-1, -0.5, 0, 0.5, 1]
    for j in range(5):
        col = "BCDEF"[j]
        put(ws, f"{col}84", f"=$B$18+({d[j]})", "f", MULT, bold=True).fill = YEAR_FILL
        put(ws, f"A{85 + j}", f"=$B$8+({d[j]})", "f", MULT, bold=True).fill = YEAR_FILL
        ws[f"A{85 + j}"].alignment = Alignment(horizontal="right")
    for i in range(5):
        r = 85 + i
        for j in range(5):
            col = "BCDEF"[j]
            fx = f'=IFERROR((($G$29*{col}$84-$G$58)/($B$7*$A{r}*(1+$B$9)+$B$10-$J$13-$J$14))^(1/5)-1,"n/a")'
            put(ws, f"{col}{r}", fx, "f", PCT, fill=BASE_FILL if (i == 2 and j == 2) else None)
    ws["A90"].value = "Leverage, interest and cash flows don't depend on the entry multiple, so only the equity cheque changes across rows."
    ws["A90"].font = SUB
    return wb


# =====================================================================
# 4. VC / PE
# =====================================================================
def vc_pe():
    wb = Workbook()
    guide(wb, "Venture Capital & Private Equity Toolkit",
          ["VC Method: works back from an exit value and a target return to the post-money and pre-money valuation, adjusted for future dilution.",
           "Cap Table: a Series A round with an option pool created pre-money (the 'option pool shuffle'), share price and ownership before and after.",
           "Exit Waterfall: payouts to Series A (1x non-participating liquidation preference), founders, seed investors and option holders at different exit values.",
           "PE Fund Waterfall: a European (whole-fund) distribution waterfall: return of capital, 8% preferred return, GP catch-up and 80/20 carried interest."],
          ["Each sheet is self-contained except the Exit Waterfall, which links (green) to the Cap Table.",
           "Change the yellow inputs and follow the formulas down the page.",
           "Then rebuild each sheet from a blank page. These four calculations come up often in VC and PE interviews and case studies."],
          ["VC Method: raise the target IRR from 40% to 50%. How much does the pre-money fall?",
           "Cap Table: change the target option pool from 15% to 20%. Who pays for the bigger pool: founders or the new investor?",
           "Exit Waterfall: find the exit value where Series A is indifferent between converting and taking its preference.",
           "Exit Waterfall: make the preference 2x. How do founders' payouts change in a small exit vs a large one?",
           "Fund Waterfall: set total proceeds to 110. Does the GP earn any carry? Why not?"])

    # --- VC method ---
    ws = wb.create_sheet("VC Method")
    title(ws, "VC Method Valuation", "($ in millions)   " + SRC)
    widths(ws, {"A": 46, "B": 14, "C": 13, "D": 13, "E": 13, "F": 13})
    section(ws, 5, "INPUTS", 2)
    for r, lbl, v, fmt in [(6, "Investment amount", 5.0, NUM), (7, "Revenue in exit year", 60, NUM),
                           (8, "Exit EV / revenue multiple", 5.0, MULT), (9, "Years to exit", 6, INT),
                           (10, "Target IRR", 0.40, PCT), (11, "Expected dilution from future rounds", 0.30, PCT)]:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", v, "key", fmt)
    section(ws, 13, "VALUATION", 2)
    calc = [
        (14, "Exit value", "=B7*B8", NUM), (15, "Required MOIC = (1 + IRR) ^ years", "=(1+B10)^B9", MULT),
        (16, "Post-money before dilution = exit value / MOIC", "=B14/B15", NUM),
        (17, "Retention ratio = 1 - future dilution", "=1-B11", PCT),
        (18, "Post-money valuation (dilution-adjusted)", "=B16*B17", NUM),
        (19, "Pre-money valuation", "=B18-B6", NUM),
        (20, "Ownership required at close", "=B6/B18", PCT), (21, "Ownership at exit (after dilution)", "=B20*B17", PCT),
        (22, "Check: value of stake at exit / investment", "=B21*B14/B6", MULT),
    ]
    for r, lbl, fx, fmt in calc:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fx, "f", fmt, bold=r in (18, 19, 20))
    ws["A23"].value = "The check should equal the required MOIC."
    ws["A23"].font = SUB
    section(ws, 25, "SENSITIVITY: PRE-MONEY VALUATION  (rows = target IRR, columns = exit EV/revenue)", 6)
    for j, d in enumerate([-2, -1, 0, 1, 2]):
        col = "BCDEF"[j]
        put(ws, f"{col}26", f"=MAX(0.5,$B$8+({d}))", "f", MULT, bold=True).fill = YEAR_FILL
    for i, d in enumerate([-0.1, -0.05, 0, 0.05, 0.1]):
        r = 27 + i
        put(ws, f"A{r}", f"=$B$10+({d})", "f", PCT, bold=True).fill = YEAR_FILL
        ws[f"A{r}"].alignment = Alignment(horizontal="right")
        for j in range(5):
            col = "BCDEF"[j]
            put(ws, f"{col}{r}", f"=$B$7*{col}$26/(1+$A{r})^$B$9*(1-$B$11)-$B$6", "f", NUM,
                fill=BASE_FILL if (i == 2 and j == 2) else None)

    # --- Cap table ---
    ws = wb.create_sheet("Cap Table")
    title(ws, "Series A Cap Table (option pool created pre-money)", "(shares in millions, $ in millions except price per share)   " + SRC)
    widths(ws, {"A": 46, "B": 14, "C": 12, "D": 14, "E": 12, "F": 15})
    section(ws, 5, "INPUTS", 2)
    for r, lbl, v, fmt in [(6, "Founder shares", 8.0, NUM), (7, "Existing option pool shares", 1.0, NUM),
                           (8, "Seed investor shares", 1.0, NUM), (9, "Series A pre-money valuation", 20, NUM),
                           (10, "Series A investment", 5, NUM), (11, "Target option pool, % of post-money", 0.15, PCT)]:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", v, "key", fmt)
    section(ws, 13, "ROUND MECHANICS", 2)
    calc = [
        (14, "Existing fully diluted shares", "=SUM(B6:B8)", NUM, None),
        (15, "Series A ownership post-money = investment / post-money", "=B10/(B9+B10)", PCT, None),
        (16, "Post-money fully diluted shares", "=(B14-B7)/(1-B11-B15)", NUM, None),
        (17, "New option pool shares (added pre-money)", "=B16*B11-B7", NUM, None),
        (18, "Price per share ($)", "=(B9+B10)/B16", USD, None),
        (19, "Series A shares issued", "=B10/B18", NUM, None),
        (20, "Effective pre-money for existing holders (price x existing shares)", "=B18*B14", NUM, None),
        (21, "Check: shares add up", '=IF(ABS(B14+B17+B19-B16)<0.0001,"OK","ERROR")', None, None),
    ]
    for r, lbl, fx, fmt, _ in calc:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fx, "f", fmt, bold=r in (18, 20))
    ws["B21"].alignment = Alignment(horizontal="right")
    ws["A22"].value = "Target pool must be larger than the existing pool. The headline pre-money includes the new pool, so existing holders effectively get the lower 'effective' pre-money."
    ws["A22"].font = SUB
    hdrs = ["Holder", "Pre-round shares", "Pre-round %", "Post-round shares", "Post-round %", "Value at post ($mm)"]
    for col, h in zip("ABCDEF", hdrs):
        ws[f"{col}24"].value = h
        ws[f"{col}24"].font = HDR
        ws[f"{col}24"].fill = HDR_FILL
        ws[f"{col}24"].alignment = Alignment(horizontal="left" if col == "A" else "right", wrap_text=True)
    holders = [(25, "Founders", "=B6", "=B6"), (26, "Seed investor", "=B8", "=B8"),
               (27, "Option pool (existing + new)", "=B7", "=B7+B17"), (28, "Series A investor", 0, "=B19")]
    for r, n, pre, post in holders:
        ws[f"A{r}"].value = n
        put(ws, f"B{r}", pre, "f", NUM)
        put(ws, f"C{r}", f"=B{r}/$B$29", "f", PCT)
        put(ws, f"D{r}", post, "f", NUM)
        put(ws, f"E{r}", f"=D{r}/$D$29", "f", PCT)
        put(ws, f"F{r}", f"=D{r}*$B$18", "f", NUM)
    ws["A29"].value = "Total"
    ws["A29"].font = BOLD
    for col, fmt in zip("BCDEF", [NUM, PCT, NUM, PCT, NUM]):
        put(ws, f"{col}29", f"=SUM({col}25:{col}28)", "f", fmt, bold=True).border = TOP
    ws["A30"].value = "Check: total value = post-money"
    put(ws, "F30", '=IF(ABS(F29-(B9+B10))<0.0001,"OK","ERROR")', "f")
    ws["F30"].alignment = Alignment(horizontal="right")

    # --- Exit waterfall ---
    ws = wb.create_sheet("Exit Waterfall")
    title(ws, "Exit Waterfall: 1x Non-Participating Preferred", "($ in millions)   " + SRC)
    widths(ws, {"A": 44, **{c: 12 for c in "BCDEFG"}})
    section(ws, 5, "INPUTS", 2)
    ws["A6"].value = "Liquidation preference multiple (x invested)"
    put(ws, "B6", 1.0, "key", MULT)
    links = [(7, "Series A invested", "='Cap Table'!B10", NUM), (8, "Series A shares", "='Cap Table'!B19", NUM),
             (9, "Total fully diluted shares", "='Cap Table'!B16", NUM), (10, "Founder shares", "='Cap Table'!B6", NUM),
             (11, "Seed shares", "='Cap Table'!B8", NUM), (12, "Option pool shares", "='Cap Table'!D27", NUM)]
    for r, lbl, fx, fmt in links:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fx, "link", fmt)
    ws["A13"].value = "Preference amount"
    put(ws, "B13", "=B6*B7", "f", NUM)
    ws["A14"].value = "Series A as-converted ownership"
    put(ws, "B14", "=B8/B9", "f", PCT)
    ws["A15"].value = "Exit value above which Series A converts"
    put(ws, "B15", "=B13/B14", "f", NUM, bold=True)
    section(ws, 17, "PAYOUTS BY EXIT VALUE", 7)
    ws["A18"].value = "Exit equity value"
    ws["A18"].font = BOLD
    rows = {19: "Series A decision", 20: "Series A payout", 21: "Left for other holders", 22: "Founders",
            23: "Seed investor", 24: "Option holders", 25: "Total distributed", 26: "Series A MOIC", 27: "Check"}
    for r, lbl in rows.items():
        ws[f"A{r}"].value = lbl
    for c, v in zip("BCDEFG", [10, 25, 50, 100, 200, 500]):
        put(ws, f"{c}18", v, "key", NUM, bold=True)
        put(ws, f"{c}19", f'=IF({c}18*$B$14>$B$13,"Convert","Preference")', "f")
        ws[f"{c}19"].alignment = Alignment(horizontal="right")
        put(ws, f"{c}20", f'=IF({c}19="Convert",{c}18*$B$14,MIN({c}18,$B$13))', "f", NUM)
        put(ws, f"{c}21", f"={c}18-{c}20", "f", NUM)
        put(ws, f"{c}22", f"={c}21*$B$10/($B$9-$B$8)", "f", NUM)
        put(ws, f"{c}23", f"={c}21*$B$11/($B$9-$B$8)", "f", NUM)
        put(ws, f"{c}24", f"={c}21*$B$12/($B$9-$B$8)", "f", NUM)
        put(ws, f"{c}25", f"={c}20+{c}22+{c}23+{c}24", "f", NUM, bold=True).border = TOP
        put(ws, f"{c}26", f"={c}20/$B$7", "f", MULT)
        put(ws, f"{c}27", f'=IF(ABS({c}25-{c}18)<0.0001,"OK","ERROR")', "f")
        ws[f"{c}27"].alignment = Alignment(horizontal="right")
    ws["A29"].value = "Simplification: all options are treated as granted and exercised, ignoring strike prices; one preferred class only."
    ws["A29"].font = SUB

    # --- PE fund waterfall ---
    ws = wb.create_sheet("PE Fund Waterfall")
    title(ws, "PE Fund Distribution Waterfall (European / whole-fund)", "($ in millions)   " + SRC)
    widths(ws, {"A": 44, "B": 14, "C": 14, "D": 14})
    section(ws, 5, "INPUTS", 2)
    for r, lbl, v, fmt in [(6, "LP capital contributed", 100, NUM), (7, "Total proceeds to distribute", 250, NUM),
                           (8, "Holding period (years, single exit)", 5, INT), (9, "Preferred return (hurdle)", 0.08, PCT),
                           (10, "Carried interest", 0.20, PCT), (11, "GP catch-up (% of distributions in that tier)", 1.0, PCT)]:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", v, "key", fmt)
    section(ws, 13, "WATERFALL", 4)
    for col, h in zip("BCD", ["Total", "To LPs", "To GP"]):
        ws[f"{col}14"].value = h
        ws[f"{col}14"].font = BOLD
        ws[f"{col}14"].alignment = Alignment(horizontal="right")
    tiers = [
        (15, "1. Return of capital", "=MIN(B7,B6)", "=B15", "=0"),
        (16, "2. Preferred return (compounded)", "=MIN(B7-B15,B6*((1+B9)^B8-1))", "=B16", "=0"),
        (17, "3. GP catch-up", "=MAX(0,MIN(B7-B15-B16,B16*B10/(B11-B10)))", "=B17*(1-B11)", "=B17*B11"),
        (18, "4. Remaining split (LP / GP)", "=B7-B15-B16-B17", "=B18*(1-B10)", "=B18*B10"),
    ]
    for r, lbl, fb, fc, fd in tiers:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fb, "f", NUM)
        put(ws, f"C{r}", fc, "f", NUM)
        put(ws, f"D{r}", fd, "f", NUM)
    ws["A19"].value = "Total"
    ws["A19"].font = BOLD
    for col in "BCD":
        put(ws, f"{col}19", f"=SUM({col}15:{col}18)", "f", NUM, bold=True).border = TOP
    ws["A20"].value = "Check: total = proceeds"
    put(ws, "B20", '=IF(ABS(B19-B7)<0.0001,"OK","ERROR")', "f")
    ws["B20"].alignment = Alignment(horizontal="right")
    section(ws, 22, "RETURNS", 4)
    for r, lbl, fx, fmt in [(23, "Gross MOIC (before carry)", "=B7/B6", MULT),
                            (24, "Gross IRR", "=B23^(1/B8)-1", PCT),
                            (25, "LP net MOIC", "=C19/B6", MULT),
                            (26, "LP net IRR", "=B25^(1/B8)-1", PCT),
                            (27, "GP share of total profit", "=IF(B7>B6,D19/(B7-B6),0)", PCT)]:
        ws[f"A{r}"].value = lbl
        put(ws, f"B{r}", fx, "f", fmt, bold=r in (25, 26))
    ws["A28"].value = "Catch-up must be larger than the carry rate. Management fees are ignored."
    ws["A28"].font = SUB
    return wb


FILES = {
    "three_statement_model.xlsx": three_statement,
    "dcf_model.xlsx": dcf,
    "lbo_model.xlsx": lbo,
    "vc_pe_toolkit.xlsx": vc_pe,
}


def build():
    os.makedirs(OUT, exist_ok=True)
    for name, fn in FILES.items():
        wb = fn()
        for ws in wb.worksheets:
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToHeight = 0
        wb.calculation.fullCalcOnLoad = True
        wb.save(os.path.join(OUT, name))
        print("wrote", name)


def embed():
    """Put the templates into index.html as base64 so the single page stays self-contained."""
    data = {}
    for name in FILES:
        with open(os.path.join(OUT, name), "rb") as fh:
            data[name] = base64.b64encode(fh.read()).decode("ascii")
    path = os.path.join(HERE, "index.html")
    with open(path, encoding="utf-8") as fh:
        html = fh.read()
    block = '<script type="application/json" id="xlsxData">' + json.dumps(data) + "</script>"
    new, n = re.subn(r'<script type="application/json" id="xlsxData">.*?</script>', lambda _: block, html, flags=re.S)
    if n != 1:
        sys.exit("xlsxData block not found in index.html")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(new)
    print("embedded", len(data), "templates into index.html")


if __name__ == "__main__":
    embed() if "--embed" in sys.argv else build()
